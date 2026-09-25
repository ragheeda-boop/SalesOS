"""Unit tests for agent_reach.contact_enrichment helpers."""

from __future__ import annotations

import csv
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import pytest

from app.modules.agent_reach.contact_enrichment import (
    EnrichmentCandidate,
    EnrichmentRow,
    QualityIssue,
    _clean_url,
    _domain_slug,
    _is_social_plausible,
    _normalize_linkedin_url,
    _normalize_whatsapp_url,
    _social_username,
    build_enrichment_rows,
    enrich_candidate,
    extract_sa_phones,
    extract_social_links,
    is_safe_company_domain,
    load_master_candidates,
    load_processed_accounts,
    normalize_domain,
    sanitize_ready_url_value,
    tier_bucket,
)
from app.modules.agent_reach.models import AgentReachResult, ChannelType

# ── normalize_domain ──────────────────────────────────────────────


def test_normalize_domain_strips_protocol():
    assert normalize_domain("https://example.com") == "example.com"


def test_normalize_domain_strips_www():
    assert normalize_domain("www.example.com") == "example.com"


def test_normalize_domain_strips_trailing_slash():
    assert normalize_domain("example.com/") == "example.com"


def test_normalize_domain_empty():
    assert normalize_domain(None) == ""
    assert normalize_domain("") == ""


def test_normalize_domain_lowercase():
    assert normalize_domain("Example.Com") == "example.com"


# ── is_safe_company_domain ────────────────────────────────────────


def test_safe_domain_basic():
    assert is_safe_company_domain("acme.com") is True


def test_safe_domain_generic_blocked():
    assert is_safe_company_domain("gmail.com") is False
    assert is_safe_company_domain("facebook.com") is False
    assert is_safe_company_domain("outlook.com") is False


def test_safe_domain_shared_threshold():
    from collections import Counter

    counts = Counter({"acme.com": 5})
    assert is_safe_company_domain("acme.com", counts) is False


def test_safe_domain_no_dot():
    assert is_safe_company_domain("localhost") is False


def test_safe_domain_empty():
    assert is_safe_company_domain("") is False


def test_safe_domain_generic_suffix():
    assert is_safe_company_domain("mail.youtube.com") is False


# ── tier_bucket ───────────────────────────────────────────────────


def test_tier_bucket_class_a():
    assert tier_bucket("CLASS A") == "CLASS A"


def test_tier_bucket_tier_b():
    assert tier_bucket("TIER B") == "TIER B"


def test_tier_bucket_tier_c():
    assert tier_bucket("TIER C") == "TIER C"
    assert tier_bucket("TIER C - INACTIVE") == "TIER C"


def test_tier_bucket_anti_icp():
    assert tier_bucket("ANTI-ICP") == "ANTI-ICP"
    assert tier_bucket("anti-icp") == "ANTI-ICP"


def test_tier_bucket_unknown():
    assert tier_bucket("") == "UNKNOWN"
    assert tier_bucket(None) == "UNKNOWN"


# ── extract_sa_phones ─────────────────────────────────────────────


def test_extract_sa_phones_international():
    phones = extract_sa_phones("Call us at +966501234567")
    assert "+966501234567" in phones


def test_extract_sa_phones_local_format():
    phones = extract_sa_phones("Call us at 0501234567")
    assert "+966501234567" in phones


def test_extract_sa_phones_short_format():
    phones = extract_sa_phones("Call us at 00966501234567")
    assert "+966501234567" in phones


def test_extract_sa_phones_toll_free():
    phones = extract_sa_phones("Toll free: 800123456")
    assert "800123456" in phones


def test_extract_sa_phones_landline():
    phones = extract_sa_phones("Landline: 0112345678")
    assert "+966112345678" in phones


def test_extract_sa_phones_deduplicates():
    phones = extract_sa_phones("+966501234567 and 0501234567")
    assert phones.count("+966501234567") == 1


def test_extract_sa_phones_empty():
    assert extract_sa_phones("") == []
    assert extract_sa_phones(None) == []


# ── extract_social_links ──────────────────────────────────────────


def test_extract_social_links_facebook():
    links = extract_social_links("Visit https://facebook.com/acme")
    assert "فيسبوك" in links
    assert any("facebook.com/acme" in u for u in links["فيسبوك"])


def test_extract_social_links_instagram():
    links = extract_social_links("Follow https://instagram.com/acme")
    assert "انستقرام" in links


def test_extract_social_links_twitter():
    links = extract_social_links("Follow https://x.com/acme")
    assert "تويتر/X" in links


def test_extract_social_links_linkedin():
    links = extract_social_links("Connect https://linkedin.com/company/acme")
    assert "لينكدإن" in links


def test_extract_social_links_rejects_linkedin_email_path():
    links = extract_social_links("Connect https://www.linkedin.com/Rak.hr@rak-j.com")
    assert "لينكدإن" not in links


def test_extract_social_links_keeps_linkedin_company_and_profile():
    company = extract_social_links("Connect https://www.linkedin.com/company/example")
    profile = extract_social_links("Connect https://www.linkedin.com/in/example-person")
    assert company["لينكدإن"] == ["https://www.linkedin.com/company/example"]
    assert profile["لينكدإن"] == ["https://www.linkedin.com/in/example-person"]


def test_extract_social_links_keeps_linkedin_school_and_showcase():
    school = extract_social_links("Connect https://www.linkedin.com/school/example")
    showcase = extract_social_links("Connect https://www.linkedin.com/showcase/example")
    assert school["لينكدإن"] == ["https://www.linkedin.com/school/example"]
    assert showcase["لينكدإن"] == ["https://www.linkedin.com/showcase/example"]


def test_extract_social_links_rejects_linkedin_invalid_first_segment():
    links = extract_social_links("Jobs https://www.linkedin.com/jobs/search/?keywords=acme")
    assert "لينكدإن" not in links


def test_normalize_linkedin_keeps_scheme_less_company_and_rejects_email_path():
    assert _normalize_linkedin_url("linkedin.com/company/example") == (
        "linkedin.com/company/example"
    )
    assert _normalize_linkedin_url("https://www.linkedin.com/Rak.hr@rak-j.com") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/shareArticle?mini=true") is None


def test_normalize_linkedin_rejects_concatenated_nested_url():
    leaked = "https://www.linkedin.com/in/https://www.linkedin.com/"
    assert _normalize_linkedin_url(leaked) is None
    links = extract_social_links(f"Connect {leaked}")
    assert "لينكدإن" not in links


def test_normalize_linkedin_rejects_platform_vendor_slug():
    assert _normalize_linkedin_url("https://www.linkedin.com/company/odoo") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/company/odoo-partner/") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/company/shopify") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/company/salla") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/company/acmethemes/") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/company/protonprivacy/") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/company/network-solutions/") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/company/grails-com") is None
    assert _normalize_linkedin_url("https://ae.linkedin.com/company/absher-business") is None
    assert _normalize_linkedin_url("https://www.linkedin.com/company/alawadhksa") == (
        "https://www.linkedin.com/company/alawadhksa"
    )
    assert _normalize_linkedin_url("https://www.linkedin.com/company/network-ksa") == (
        "https://www.linkedin.com/company/network-ksa"
    )


def test_extract_social_links_rejects_odoo_footer_linkedin():
    text = (
        "Powered by Odoo https://www.linkedin.com/company/odoo "
        "and us https://www.linkedin.com/company/alawadhksa"
    )
    links = extract_social_links(text, company_domain="alawadhksa.com")
    assert links.get("لينكدإن") == ["https://www.linkedin.com/company/alawadhksa"]


def test_extract_social_links_rejects_acmethemes_footer_linkedin():
    text = (
        "Theme by ACME https://www.linkedin.com/company/acmethemes/ "
        "and us https://www.linkedin.com/company/aicindus"
    )
    links = extract_social_links(text, company_domain="aicindus.com")
    assert links.get("لينكدإن") == ["https://www.linkedin.com/company/aicindus"]


def test_extract_social_links_rejects_grails_footer_linkedin():
    text = (
        "Built with Grails https://www.linkedin.com/company/grails-com "
        "and us https://www.linkedin.com/company/safa-ksa"
    )
    links = extract_social_links(text, company_domain="safa.com")
    assert links.get("لينكدإن") == ["https://www.linkedin.com/company/safa-ksa"]


def test_extract_social_links_rejects_protonprivacy_footer_linkedin():
    text = (
        "Powered by Proton https://www.linkedin.com/company/protonprivacy/ "
        "and us https://www.linkedin.com/company/passinbox"
    )
    links = extract_social_links(text, company_domain="passinbox.com")
    assert links.get("لينكدإن") == ["https://www.linkedin.com/company/passinbox"]


def test_extract_social_links_rejects_network_solutions_footer_linkedin():
    text = (
        "Hosted by iPage https://www.linkedin.com/company/network-solutions/ "
        "and us https://www.linkedin.com/company/emailmg"
    )
    links = extract_social_links(text, company_domain="emailmg.ipage.com")
    assert links.get("لينكدإن") == ["https://www.linkedin.com/company/emailmg"]


def test_extract_social_links_rejects_concatenated_http_urls():
    links = extract_social_links(
        "Chat https://api.whatsapp.com/send?phone=966500000000&text=https://example.com/"
    )
    assert "واتساب" not in links


def test_clean_url_strips_trailing_middle_dot_garbage():
    raw = "https://twitter.com/accbs_sa)·"
    assert _clean_url(raw) == "https://twitter.com/accbs_sa"


def test_clean_url_strips_trailing_markdown_hashes():
    raw = (
        "https://www.linkedin.com/in/"
        "%D8%A2%D9%84-%D8%B7%D8%A7%D9%88%D9%89-128239373/)###"
    )
    assert _clean_url(raw).endswith("128239373/")
    assert ")###" not in _clean_url(raw)


def test_normalize_linkedin_rejects_extra_ui_path():
    dirty = "https://www.linkedin.com/company/ucc-holding/mycompany/verification/"
    assert _normalize_linkedin_url(dirty) is None
    links = extract_social_links(f"Connect {dirty}")
    assert "لينكدإن" not in links


def test_extract_social_links_whatsapp_sa():
    links = extract_social_links("WhatsApp: https://wa.me/966501234567")
    assert "واتساب" in links


def test_extract_social_links_whatsapp_non_sa_ignored():
    links = extract_social_links("WhatsApp: https://wa.me/12025551234")
    assert "واتساب" not in links


def test_extract_social_links_empty():
    assert extract_social_links("") == {}


def test_extract_social_links_bare_homepage_filtered():
    links = extract_social_links("Visit https://www.facebook.com/ or https://twitter.com/")
    assert "فيسبوك" not in links
    assert "تويتر/X" not in links


def test_extract_social_links_trailing_artifacts_cleaned():
    links = extract_social_links("https://www.instagram.com/acme)[")
    assert "انستقرام" in links
    assert links["انستقرام"][0] == "https://www.instagram.com/acme"


def test_extract_social_links_whatsapp_strips_markdown_image_artifact():
    """Confirmed leak: markdown )![Image / ![Image must not stay on WhatsApp URLs."""
    api_links = extract_social_links(
        "Chat https://api.whatsapp.com/send?phone=966920012777&text=)![Image leftover"
    )
    me_links = extract_social_links("Chat https://wa.me/966566222773)![Image leftover")
    non_sa_links = extract_social_links("Chat https://wa.me/12025551234)![Image leftover")

    assert "واتساب" in api_links
    assert api_links["واتساب"][0] == (
        "https://api.whatsapp.com/send?phone=966920012777"
    )
    assert "![Image" not in api_links["واتساب"][0]
    assert "text=" not in api_links["واتساب"][0]

    assert "واتساب" in me_links
    assert me_links["واتساب"][0] == "https://wa.me/966566222773"
    assert "![Image" not in me_links["واتساب"][0]

    assert "واتساب" not in non_sa_links


def test_normalize_whatsapp_strips_leftover_text_query():
    """Confirmed remainder leak: &text=)Contact must not stay on WhatsApp URLs."""
    leaked = "https://web.whatsapp.com/send?phone=966535055057&text=)Contact"
    assert _normalize_whatsapp_url(leaked) == (
        "https://web.whatsapp.com/send?phone=966535055057"
    )
    links = extract_social_links(f"Chat {leaked}")
    assert links["واتساب"] == ["https://web.whatsapp.com/send?phone=966535055057"]
    assert "text=" not in links["واتساب"][0]
    assert ")Contact" not in links["واتساب"][0]


def test_extract_social_links_deduplication():
    text = "https://facebook.com/a https://facebook.com/a"
    links = extract_social_links(text)
    assert len(links.get("فيسبوك", [])) == 1


# ── build_enrichment_rows ─────────────────────────────────────────


def _make_candidate(**overrides) -> EnrichmentCandidate:
    defaults = {
        "account_id": "ACC-001",
        "company_name": "Acme Corp",
        "domain": "acme.com",
        "tier": "TIER C",
    }
    defaults.update(overrides)
    return EnrichmentCandidate(**defaults)


def test_build_enrichment_rows_phone_and_social():
    page = "Call +966501234567 Visit https://facebook.com/acme"
    cand = _make_candidate()
    rows = build_enrichment_rows(cand, page)
    fields = [r.field for r in rows]
    assert "جوال" in fields
    assert "فيسبوك" in fields


def test_build_enrichment_rows_skip_existing_phone():
    page = "Call +966501234567"
    cand = _make_candidate(has_phone=True)
    rows = build_enrichment_rows(cand, page)
    assert all(r.field != "جوال" for r in rows)


def test_build_enrichment_rows_skip_existing_social():
    page = "Visit https://facebook.com/acme"
    cand = _make_candidate(has_social=True)
    rows = build_enrichment_rows(cand, page)
    assert all(r.field != "فيسبوك" for r in rows)


def test_build_enrichment_rows_empty_page():
    cand = _make_candidate()
    rows = build_enrichment_rows(cand, "")
    assert rows == []


# ── EnrichmentRow.as_csv_row ─────────────────────────────────────


def test_enrichment_row_as_csv():
    row = EnrichmentRow("A", "B", "C", "D", "E", "F")
    assert row.as_csv_row() == ["A", "B", "C", "D", "E", "F"]


# ── QualityIssue.as_csv_row ───────────────────────────────────────


def test_quality_issue_as_csv():
    issue = QualityIssue("A", "B", "C", "problem")
    assert issue.as_csv_row() == ["A", "B", "C", "problem"]


# ── load_master_candidates ────────────────────────────────────────


def _write_master_csv(rows: list[dict]) -> Path:
    tmpdir = Path(tempfile.mkdtemp())
    csv_path = tmpdir / "master.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "Master Account ID",
                "Canonical_Company_Name",
                "Primary_Domain",
                "Account_Tier_v2",
                "City",
                "has_phone",
                "has_social",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)
    return csv_path


def test_load_master_candidates_filters_processed():
    csv_path = _write_master_csv(
        [
            {
                "Master Account ID": "A1",
                "Canonical_Company_Name": "Acme",
                "Primary_Domain": "acme.com",
                "Account_Tier_v2": "TIER C",
                "City": "Riyadh",
                "has_phone": "0",
                "has_social": "0",
            },
            {
                "Master Account ID": "A2",
                "Canonical_Company_Name": "Beta",
                "Primary_Domain": "beta.com",
                "Account_Tier_v2": "TIER C",
                "City": "Jeddah",
                "has_phone": "0",
                "has_social": "0",
            },
        ]
    )
    candidates = load_master_candidates(csv_path, processed_accounts={"A1"})
    assert len(candidates) == 1
    assert candidates[0].account_id == "A2"


def test_load_master_candidates_skips_generic_domain():
    csv_path = _write_master_csv(
        [
            {
                "Master Account ID": "A1",
                "Canonical_Company_Name": "NoDomain",
                "Primary_Domain": "gmail.com",
                "Account_Tier_v2": "TIER C",
                "City": "Riyadh",
                "has_phone": "0",
                "has_social": "0",
            },
        ]
    )
    candidates = load_master_candidates(csv_path)
    assert len(candidates) == 0


def test_load_master_candidates_skips_existing_phone_and_social():
    csv_path = _write_master_csv(
        [
            {
                "Master Account ID": "A1",
                "Canonical_Company_Name": "Complete",
                "Primary_Domain": "complete.com",
                "Account_Tier_v2": "TIER C",
                "City": "Riyadh",
                "has_phone": "1",
                "has_social": "1",
            },
        ]
    )
    candidates = load_master_candidates(csv_path)
    assert len(candidates) == 0


def test_load_master_candidates_tier_filter():
    csv_path = _write_master_csv(
        [
            {
                "Master Account ID": "A1",
                "Canonical_Company_Name": "TierB",
                "Primary_Domain": "tierb.com",
                "Account_Tier_v2": "TIER B",
                "City": "",
                "has_phone": "0",
                "has_social": "0",
            },
            {
                "Master Account ID": "A2",
                "Canonical_Company_Name": "TierC",
                "Primary_Domain": "tierc.com",
                "Account_Tier_v2": "TIER C",
                "City": "",
                "has_phone": "0",
                "has_social": "0",
            },
        ]
    )
    candidates = load_master_candidates(csv_path, tiers={"TIER C"})
    assert len(candidates) == 1
    assert candidates[0].account_id == "A2"


# ── load_processed_accounts ───────────────────────────────────────


def test_load_processed_accounts_from_workbook():
    from openpyxl import Workbook

    tmpdir = Path(tempfile.mkdtemp())
    xlsx_path = tmpdir / "test.xlsx"
    wb = Workbook()
    ws = wb.create_sheet("مرشحون جاهزون")
    ws.append(["رقم الحساب", "اسم الشركة"])
    ws.append(["ACC-001", "Acme"])
    ws.append(["ACC-002", "Beta"])
    wb.save(xlsx_path)

    processed = load_processed_accounts(xlsx_path)
    assert "ACC-001" in processed
    assert "ACC-002" in processed


def test_load_processed_accounts_missing_sheet():
    from openpyxl import Workbook

    tmpdir = Path(tempfile.mkdtemp())
    xlsx_path = tmpdir / "empty.xlsx"
    wb = Workbook()
    wb.save(xlsx_path)

    processed = load_processed_accounts(xlsx_path)
    assert len(processed) == 0


# ── enrich_candidate (mocked) ────────────────────────────────────


@pytest.mark.asyncio
async def test_enrich_candidate_success():
    candidate = _make_candidate(domain="acme.com")
    mock_service = MagicMock()
    page_content = (
        "Welcome to Acme Corp. Call us at +966501234567 for inquiries. "
        "Follow us on social media: https://facebook.com/acme and "
        "https://instagram.com/acme. We provide great services."
    )
    mock_service.read_web_page = AsyncMock(
        return_value=AgentReachResult(
            success=True,
            channel=ChannelType.WEB,
            action="read",
            data=page_content,
        )
    )
    rows, issue = await enrich_candidate(candidate, mock_service)
    assert len(rows) > 0
    assert issue is None


@pytest.mark.asyncio
async def test_enrich_candidate_read_failure():
    candidate = _make_candidate(domain="acme.com")
    mock_service = MagicMock()
    mock_service.read_web_page = AsyncMock(
        return_value=AgentReachResult(
            success=False,
            channel=ChannelType.WEB,
            action="read",
            error="Connection refused",
        )
    )
    rows, issue = await enrich_candidate(candidate, mock_service)
    assert rows == []
    assert issue is not None
    assert "Agent Reach" in issue.issue


@pytest.mark.asyncio
async def test_enrich_candidate_no_data():
    candidate = _make_candidate(domain="acme.com")
    mock_service = MagicMock()
    mock_service.read_web_page = AsyncMock(
        return_value=AgentReachResult(
            success=True,
            channel=ChannelType.WEB,
            action="read",
            data=None,
        )
    )
    rows, issue = await enrich_candidate(candidate, mock_service)
    assert rows == []
    assert issue is not None


@pytest.mark.asyncio
async def test_enrich_candidate_short_content():
    candidate = _make_candidate(domain="acme.com")
    mock_service = MagicMock()
    mock_service.read_web_page = AsyncMock(
        return_value=AgentReachResult(
            success=True,
            channel=ChannelType.WEB,
            action="read",
            data="short",
        )
    )
    rows, issue = await enrich_candidate(candidate, mock_service)
    assert rows == []
    assert issue is not None
    assert "قصير" in issue.issue


@pytest.mark.asyncio
async def test_enrich_candidate_no_usable_data():
    candidate = _make_candidate(domain="acme.com")
    mock_service = MagicMock()
    mock_service.read_web_page = AsyncMock(
        return_value=AgentReachResult(
            success=True,
            channel=ChannelType.WEB,
            action="read",
            data="This is a long page about corporate governance with no contact info at all. " * 3,
        )
    )
    rows, issue = await enrich_candidate(candidate, mock_service)
    assert rows == []
    assert issue is not None
    assert "لم يعثر" in issue.issue


@pytest.mark.asyncio
async def test_enrich_candidate_invalid_ipv6_does_not_raise():
    candidate = _make_candidate(domain="ajaz.com.sa", has_phone=True)
    mock_service = MagicMock()
    page_content = (
        " millennial page text with broken host https://[unclosed-ipv6 "
        "and a clean company page https://linkedin.com/company/ajaz "
        + ("content " * 30)
    )
    mock_service.read_web_page = AsyncMock(
        return_value=AgentReachResult(
            success=True,
            channel=ChannelType.WEB,
            action="read",
            data=page_content,
        )
    )
    rows, issue = await enrich_candidate(candidate, mock_service)
    assert issue is None
    assert any(row.field == "لينكدإن" for row in rows)


# ── _domain_slug ──────────────────────────────────────────────────


def test_domain_slug_basic():
    assert _domain_slug("acme.com") == "acme"


def test_domain_slug_hyphenated():
    assert _domain_slug("the-carbon-steel.com") == "thecarbonsteel"


def test_domain_slug_subdomain():
    assert _domain_slug("en-sa.ajmal.com") == "ensa"


# ── _social_username ──────────────────────────────────────────────


def test_social_username_twitter():
    assert _social_username("https://x.com/alarmlawfirm") == "alarmlawfirm"


def test_social_username_instagram():
    assert _social_username("https://www.instagram.com/ajmalsaudi/") == "ajmalsaudi"


def test_social_username_facebook_page():
    assert _social_username("https://www.facebook.com/AjmalSaudiPerfumes/") == "ajmalsaudiperfumes"


def test_social_username_empty_path():
    assert _social_username("https://www.facebook.com/") == ""


# ── _is_social_plausible ─────────────────────────────────────────


def test_is_social_plausible_exact():
    assert _is_social_plausible("acme", "acme") is True


def test_is_social_plausible_contains():
    assert _is_social_plausible("acmecorp", "acme") is True


def test_is_social_plausible_reverse():
    assert _is_social_plausible("acme", "acmecorp") is True


def test_is_social_plausible_unrelated():
    assert _is_social_plausible("obastidortv", "law1s") is False


def test_is_social_plausible_empty():
    assert _is_social_plausible("", "acme") is False


def test_is_social_plausible_rejects_single_letter_substring():
    assert _is_social_plausible("p", "cityplaza") is False
    assert _is_social_plausible("acme", "") is False


# ── extract_social_links plausibility ─────────────────────────────


def test_extract_social_links_rejects_unrelated_handle():
    text = "Follow https://x.com/obastidortv and https://x.com/acme"
    links = extract_social_links(text, company_domain="acme.com")
    assert "تويتر/X" in links
    assert len(links["تويتر/X"]) == 1
    assert "acme" in links["تويتر/X"][0]


def test_extract_social_links_accepts_related_handle():
    text = "Follow https://x.com/alarmlawfirm"
    links = extract_social_links(text, company_domain="alarmlawfirm.com")
    assert "تويتر/X" in links


def test_extract_social_links_no_domain_skips_check():
    text = "Follow https://x.com/anything"
    links = extract_social_links(text)
    assert "تويتر/X" in links


def test_extract_social_links_whatsapp_skips_plausibility():
    text = "WhatsApp: https://wa.me/966501234567"
    links = extract_social_links(text, company_domain="acme.com")
    assert "واتساب" in links


def test_extract_social_links_rejects_tracking_subdomain():
    text = "Tracker: https://analytics.twitter.com/i/adsct?foo=bar"
    links = extract_social_links(text, company_domain="acme.com")
    assert "تويتر/X" not in links


def test_extract_social_links_rejects_tweet_status_url():
    text = "Tweet: https://twitter.com/acme/status/123456"
    links = extract_social_links(text, company_domain="acme.com")
    assert "تويتر/X" not in links


def test_extract_social_links_rejects_instagram_post_and_reel():
    posts = extract_social_links(
        "https://www.instagram.com/p/DbvHvWsn5n0/ and https://www.instagram.com/reel/abc123/",
        company_domain="cityplaza.com",
    )
    assert "انستقرام" not in posts


def test_extract_social_links_rejects_twitter_internal_path():
    text = "Link: https://x.com/i/jf/stories/home"
    links = extract_social_links(text, company_domain="acme.com")
    assert "تويتر/X" not in links


def test_extract_social_links_rejects_twitter_search():
    text = "Link: https://twitter.com/search?q=acme"
    links = extract_social_links(text, company_domain="acme.com")
    assert "تويتر/X" not in links


def test_extract_social_links_strips_twitter_tracking_query():
    """Share-tracking ?t= / &s=08 must not stay on a genuine X profile."""
    text = "Follow https://x.com/SnoodHotels?t=eCw4g1A0-tXCoVk6ghOgsg&s=08"
    links = extract_social_links(text, company_domain="snoodhotels.com")
    assert "تويتر/X" in links
    assert links["تويتر/X"] == ["https://x.com/SnoodHotels"]
    assert "?t=" not in links["تويتر/X"][0]
    assert "s=08" not in links["تويتر/X"][0]


def test_extract_social_links_rejects_twitter_intent_url():
    tweet = extract_social_links("Share https://twitter.com/intent/tweet?text=hello")
    follow = extract_social_links("Follow https://x.com/intent/follow?screen_name=acme")
    assert "تويتر/X" not in tweet
    assert "تويتر/X" not in follow


def test_extract_social_links_rejects_twitter_share_path():
    links = extract_social_links(
        "Share https://x.com/share?url=https://example.com/page"
    )
    assert "تويتر/X" not in links


def test_extract_social_links_keeps_clean_x_and_twitter_handles():
    x_links = extract_social_links("Follow https://x.com/SnoodHotels")
    tw_links = extract_social_links("Follow https://twitter.com/SnoodHotels")
    assert x_links["تويتر/X"] == ["https://x.com/SnoodHotels"]
    assert tw_links["تويتر/X"] == ["https://twitter.com/SnoodHotels"]


def test_extract_social_links_ignores_invalid_ipv6_url():
    """Unclosed IPv6 hosts must not crash extraction (cycle-25 / retry-14)."""
    text = (
        "Broken https://[unclosed-ipv6 leftover and "
        "https://instagram.com/acme plus https://facebook.com/acme"
    )
    links = extract_social_links(text)
    assert links["انستقرام"][0] == "https://instagram.com/acme"
    assert links["فيسبوك"][0] == "https://facebook.com/acme"


def test_build_enrichment_rows_survives_invalid_ipv6_url():
    candidate = _make_candidate(has_phone=True, has_social=False)
    page = (
        "Company page with a broken markdown host https://[::1 "
        "and a usable profile https://linkedin.com/company/acme "
        + ("more text " * 20)
    )
    rows = build_enrichment_rows(candidate, page)
    assert any(row.field == "لينكدإن" and "linkedin.com/company/acme" in row.value for row in rows)


# ── _clean_url ────────────────────────────────────────────────────


def test_clean_url_strips_markdown_image_suffix():
    raw = "https://wa.me/966551234567)![Image"
    assert _clean_url(raw) == "https://wa.me/966551234567"


def test_clean_url_strips_trailing_parens_and_bang():
    raw = "https://wa.me/966551234567)!"
    assert _clean_url(raw) == "https://wa.me/966551234567"


def test_clean_url_strips_trailing_markdown_asterisks():
    raw = "https://www.tiktok.com/@flashmed.ksa)**"
    assert _clean_url(raw) == "https://www.tiktok.com/@flashmed.ksa"


def test_extract_social_links_strips_tiktok_markdown_asterisks():
    links = extract_social_links(
        "Follow https://www.tiktok.com/@flashmed.ksa)** leftover",
        company_domain="flashmed.net",
    )
    assert links["تيك توك"] == ["https://www.tiktok.com/@flashmed.ksa"]


def test_clean_url_strips_confirmed_workbook_leaks():
    api = "https://api.whatsapp.com/send?phone=966920012777&text=)![Image"
    me = "https://wa.me/966566222773)![Image"
    assert _clean_url(api) == "https://api.whatsapp.com/send?phone=966920012777&text="
    assert _clean_url(me) == "https://wa.me/966566222773"


def test_sanitize_ready_url_value_strips_image_artifact_only():
    api = "https://api.whatsapp.com/send?phone=966920012777&text=)![Image"
    me = "https://wa.me/966566222773)![Image"
    assert sanitize_ready_url_value(api) == (
        "https://api.whatsapp.com/send?phone=966920012777&text="
    )
    assert sanitize_ready_url_value(me) == "https://wa.me/966566222773"
    assert "![Image" not in sanitize_ready_url_value(api)
    assert "![Image" not in sanitize_ready_url_value(me)


def test_sanitize_ready_url_value_leaves_non_artifact_ready_values():
    arabic_linkedin = (
        "https://www.linkedin.com/company/"
        "شركة-عبدالرحمن-عبدالعزيز-السديس-للخدمات-اللوجستية"
    )
    dotted_linkedin = (
        "https://www.linkedin.com/company/madaf-trading-and-contracting-co.-ltd."
    )
    phone = "+966501234567"
    assert sanitize_ready_url_value(arabic_linkedin) == arabic_linkedin
    assert sanitize_ready_url_value(dotted_linkedin) == dotted_linkedin
    assert sanitize_ready_url_value(phone) == phone
