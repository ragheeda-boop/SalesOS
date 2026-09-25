"""Unit tests for the review-only Agent Reach master enrichment patch."""

from __future__ import annotations

import csv
import shutil
import sys
import tempfile
from pathlib import Path

import pytest

BACKEND_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = BACKEND_ROOT / "scripts"
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from agent_reach_master_enrichment_patch import (  # noqa: E402
    ADDITIVE_COLUMNS,
    PATCH_HEADERS,
    QUALITY_HEADERS,
    STATUS_QUALITY,
    STATUS_READY,
    STATUS_READY_WITH_ISSUES,
    VALUE_DELIMITER,
    IssueRecord,
    ReadyRecord,
    additive_columns_for,
    build_account_patches,
    build_patch_rows,
    build_quality_rows,
    classify_quality_issue,
    discover_master_id_column,
    enrichment_status,
    field_to_column,
    join_values,
    normalize_patch_value,
    quality_flag,
    recover_linkedin_value,
    recover_phone_values,
    recover_ready_values,
    recover_twitter_value,
    recover_whatsapp_value,
    run_patch,
    write_enriched_master,
)

from app.modules.agent_reach.contact_enrichment import ISSUES_SHEET, READY_SHEET  # noqa: E402


def _ready(account_id: str, field: str, value: str) -> ReadyRecord:
    return ReadyRecord(account_id, "Acme", "acme.com", "TIER C", field, value)


def _issue(
    account_id: str,
    text: str,
    *,
    company: str = "Acme",
    domain: str = "acme.com",
) -> IssueRecord:
    return IssueRecord(account_id, company, domain, text)


def _write_master(path: Path, rows: list[dict[str, str]]) -> Path:
    fieldnames = [
        "Master Account ID",
        "CR_Numbers",
        "Primary_Domain",
        "Primary_Phone",
        "Primary_Email",
        "Canonical_Company_Name",
        "Account_Tier_v2",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return path


def _write_workbook(path: Path, ready: list[ReadyRecord], issues: list[IssueRecord]) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ready_ws = wb.create_sheet(READY_SHEET)
    ready_ws.append(["رقم الحساب", "اسم الشركة", "النطاق", "الفئة", "الحقل", "القيمة/الرابط"])
    for row in ready:
        ready_ws.append(
            [row.account_id, row.company_name, row.domain, row.tier, row.field, row.value]
        )
    issue_ws = wb.create_sheet(ISSUES_SHEET)
    issue_ws.append(["رقم الحساب", "اسم الشركة", "النطاق المسجل", "المشكلة"])
    for row in issues:
        issue_ws.append([row.account_id, row.company_name, row.domain, row.issue])
    default = wb[wb.sheetnames[0]]
    if default.title not in {READY_SHEET, ISSUES_SHEET}:
        wb.remove(default)
    wb.save(path)
    return path


def test_discover_master_id_column():
    assert discover_master_id_column(["CR_Numbers", "Master Account ID"]) == "Master Account ID"
    assert discover_master_id_column(["master_account_id"]) == "master_account_id"


def test_arabic_field_mapping():
    assert field_to_column("جوال") == "AgentReach_Phones"
    assert field_to_column("هاتف") == "AgentReach_Phones"
    assert field_to_column("واتساب") == "AgentReach_WhatsApp"
    assert field_to_column("لينكدإن") == "AgentReach_LinkedIn"
    assert field_to_column("إنستغرام") == "AgentReach_Instagram"
    assert field_to_column("انستقرام") == "AgentReach_Instagram"
    assert field_to_column("تويتر/X") == "AgentReach_X"
    assert field_to_column("فيسبوك") == "AgentReach_Facebook"
    assert field_to_column("تيك توك") == "AgentReach_TikTok"
    assert field_to_column("يوتيوب") == "AgentReach_YouTube"
    assert field_to_column("سناب") == "AgentReach_Snapchat"
    assert field_to_column("unknown") is None


def test_status_mapping():
    assert STATUS_READY == "PENDING_REVIEW"
    assert STATUS_QUALITY == "QUALITY_BACKLOG"
    assert enrichment_status(True, True) == STATUS_READY_WITH_ISSUES
    assert enrichment_status(True, False) == STATUS_READY
    assert enrichment_status(False, True) == STATUS_QUALITY
    assert enrichment_status(False, False) == ""
    assert quality_flag(STATUS_READY) == "ok"
    assert quality_flag(STATUS_READY_WITH_ISSUES) == "needs_review"
    assert quality_flag(STATUS_QUALITY) == "needs_review"


def test_delimiter_has_no_spaces():
    assert VALUE_DELIMITER == "|"
    assert join_values(["+966501111111", "+966502222222"]) == "+966501111111|+966502222222"
    assert " | " not in join_values(["a", "b"])


def test_dedupe_phones_and_urls():
    patches = build_account_patches(
        [
            _ready("MA-1", "جوال", "0501234567"),
            _ready("MA-1", "جوال", "+966501234567"),
            _ready("MA-1", "جوال", "966501234567"),
            _ready("MA-1", "لينكدإن", "https://www.linkedin.com/company/acme/"),
            _ready("MA-1", "لينكدإن", "https://www.linkedin.com/company/acme"),
            _ready("MA-1", "واتساب", "https://wa.me/966501234567)![Image"),
        ],
        [],
    )
    patch = patches["MA-1"]
    assert patch.values_by_column["AgentReach_Phones"] == ["+966501234567"]
    assert patch.values_by_column["AgentReach_LinkedIn"] == [
        "https://www.linkedin.com/company/acme/"
    ]
    assert patch.values_by_column["AgentReach_WhatsApp"] == ["https://wa.me/966501234567"]
    assert patch.duplicates_removed > 0


def test_whatsapp_and_phone_normalization():
    assert normalize_patch_value("جوال", "0551234567") == "+966551234567"
    assert normalize_patch_value("واتساب", "https://wa.me/966551234567&text=)Contact") == (
        "https://wa.me/966551234567"
    )


def test_quality_issues_stay_out_of_social_columns():
    patches = build_account_patches(
        [_ready("MA-2", "فيسبوك", "https://facebook.com/acme")],
        [_issue("MA-2", "النطاق لا يعمل (فشل DNS)")],
    )
    additive = additive_columns_for(patches["MA-2"])
    assert additive["AgentReach_Facebook"] == "https://facebook.com/acme"
    assert "فشل DNS" not in additive["AgentReach_Facebook"]
    assert "فشل DNS" in additive["AgentReach_Quality_Issue"]
    assert additive["AgentReach_Enrichment_Status"] == STATUS_READY_WITH_ISSUES
    assert additive["Domain_Quality_Status"] == "UNREACHABLE"


@pytest.mark.parametrize(
    ("text", "category", "status"),
    [
        ("النطاق المسجل يبدو خطأ إملائياً/نطاقاً غير رسمي", "typo_domain", "TYPO_DOMAIN"),
        ("النطاق لا يعمل، تعذر الاتصال بالخادم (timeout)", "timeout", "TIMEOUT"),
        ("قرأ Agent Reach الموقع لكن لم يعثر على جوال", "no_contact_found", "NO_CONTACT_FOUND"),
        (
            "الموقع يمثل شركة مختلفة تماماً عن اسم الحساب",
            "company_domain_mismatch",
            "COMPANY_DOMAIN_MISMATCH",
        ),
        ("نطاق حكومي (.gov.sa) موجود خطأً", "government_domain", "GOVERNMENT_DOMAIN"),
    ],
)
def test_classify_quality_issue(text: str, category: str, status: str):
    got_category, got_status, action = classify_quality_issue(text)
    assert got_category == category
    assert got_status == status
    assert action


def test_unmatched_ids_do_not_invent_master_rows(tmp_path: Path):
    master = _write_master(
        tmp_path / "master.csv",
        [
            {
                "Master Account ID": "MA-100",
                "CR_Numbers": "1010",
                "Primary_Domain": "kept.com",
                "Primary_Phone": "011111",
                "Primary_Email": "a@kept.com",
                "Canonical_Company_Name": "Kept Co",
                "Account_Tier_v2": "TIER C",
            }
        ],
    )
    patches = build_account_patches(
        [
            _ready("MA-100", "جوال", "+966501111111"),
            ReadyRecord("MA-999", "Ghost", "acme.com", "TIER C", "جوال", "+966509999999"),
        ],
        [_issue("MA-888", "timeout", company="Missing")],
    )
    enriched = tmp_path / "enriched.csv"
    stats = write_enriched_master(master, enriched, patches)
    with enriched.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    ids = [row["Master Account ID"] for row in rows]
    assert ids == ["MA-100"]
    assert "MA-999" not in ids
    assert "MA-888" not in ids
    assert patches["MA-999"].unmatched is True
    assert patches["MA-888"].unmatched is True
    assert "MA-100" in stats["matched_ids"]

    patch_rows = build_patch_rows(patches, master_names=stats["master_names"])
    quality_rows = build_quality_rows(patches)
    patch_ids = {row[0] for row in patch_rows}
    quality_ids = {row[0] for row in quality_rows}
    assert "MA-999" in patch_ids
    assert "MA-888" in quality_ids
    ghost = next(row for row in patch_rows if row[0] == "MA-999")
    assert "unmatched_master_id" in ghost[-1]


def test_original_columns_unchanged(tmp_path: Path):
    original = {
        "Master Account ID": "MA-100",
        "CR_Numbers": "1005; 7066",
        "Primary_Domain": "original.com",
        "Primary_Phone": "0112345678",
        "Primary_Email": "keep@original.com",
        "Canonical_Company_Name": "Original Name",
        "Account_Tier_v2": "CLASS A",
    }
    master = _write_master(tmp_path / "master.csv", [original])
    patches = build_account_patches(
        [
            _ready("MA-100", "جوال", "+966501234567"),
            _ready("MA-100", "انستقرام", "https://instagram.com/acme"),
        ],
        [],
    )
    enriched = tmp_path / "enriched.csv"
    write_enriched_master(master, enriched, patches)
    with enriched.open("r", encoding="utf-8-sig", newline="") as handle:
        row = next(csv.DictReader(handle))
    for key, value in original.items():
        assert row[key] == value
    assert row["AgentReach_Phones"] == "+966501234567"
    assert row["AgentReach_Instagram"] == "https://instagram.com/acme"
    assert row["AgentReach_Enrichment_Status"] == STATUS_READY
    assert row["AgentReach_Ready_Row_Count"] == "2"
    for column in ADDITIVE_COLUMNS:
        assert column in row


def test_run_patch_writes_only_outputs(tmp_path: Path):
    output_root = BACKEND_ROOT / "outputs" / "agent_reach_contact_enrichment"
    work_dir = Path(tempfile.mkdtemp(prefix="patch_test_", dir=output_root))
    master = _write_master(
        tmp_path / "master.csv",
        [
            {
                "Master Account ID": "MA-1",
                "CR_Numbers": "",
                "Primary_Domain": "acme.com",
                "Primary_Phone": "",
                "Primary_Email": "",
                "Canonical_Company_Name": "Acme Co",
                "Account_Tier_v2": "TIER C",
            }
        ],
    )
    master_stat = master.stat()
    workbook = _write_workbook(
        tmp_path / "wb.xlsx",
        [
            _ready("MA-1", "جوال", "0501112233"),
            _ready("MA-1", "جوال", "+966501112233"),
            _ready("MA-2", "فيسبوك", "https://facebook.com/ghost"),
        ],
        [_issue("MA-1", "الموقع يمثل شركة مختلفة تماماً عن اسم الحساب")],
    )
    report_path = work_dir / "MASTER_ENRICHMENT_PATCH_REPORT.md"
    try:
        stats = run_patch(
            workbook_path=workbook,
            master_csv_path=master,
            output_dir=work_dir,
            report_path=report_path,
        )
        after = master.stat()
        assert after.st_size == master_stat.st_size
        assert after.st_mtime == master_stat.st_mtime
        assert stats["master_unchanged"] is True
        assert stats["outputs_only"] is True
        assert stats["total_master_rows"] == 1
        assert stats["matched_account_ids"] == 1
        assert stats["unmatched_ids"] == ["MA-2"]
        assert stats["unique_accounts_enriched"] == 1
        assert stats["ready_rows_applied"] == 1
        assert stats["duplicates_removed"] == 1
        assert stats["quality_issue_count"] == 1

        patch_csv = work_dir / "master_enrichment_patch.csv"
        quality_csv = work_dir / "master_quality_issues_patch.csv"
        enriched_csv = work_dir / "01_Master_Accounts_enriched_review_only.csv"
        with patch_csv.open("r", encoding="utf-8-sig", newline="") as handle:
            patch_reader = csv.DictReader(handle)
            assert patch_reader.fieldnames == PATCH_HEADERS
            patch_rows = list(patch_reader)
        assert patch_rows[0]["review_status"] == STATUS_READY_WITH_ISSUES
        assert patch_rows[0]["source_method"] == "agent_reach_review_workbook"
        with quality_csv.open("r", encoding="utf-8-sig", newline="") as handle:
            quality_reader = csv.DictReader(handle)
            assert quality_reader.fieldnames == QUALITY_HEADERS
            quality_rows = list(quality_reader)
        assert quality_rows[0]["issue_category"] == "company_domain_mismatch"
        with enriched_csv.open("r", encoding="utf-8-sig", newline="") as handle:
            enriched = next(csv.DictReader(handle))
        assert enriched["Primary_Domain"] == "acme.com"
        assert enriched["AgentReach_Phones"] == "+966501112233"
        assert enriched["AgentReach_Enrichment_Status"] == STATUS_READY_WITH_ISSUES
        text = report_path.read_text(encoding="utf-8")
        assert "REVIEW-ONLY MASTER PATCH CREATED" in text
        assert "NO ORIGINAL MASTER OVERWRITE" in text
        assert "PRODUCTION NOT APPROVED" in text
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_recover_drops_false_positives_and_keeps_saudi_values():
    assert recover_phone_values("0112782590 / 0508040404") == [
        "+966508040404",
        "+966112782590",
    ]
    assert recover_phone_values("+996572751571") == ["+966572751571"]
    assert recover_phone_values("00112710231") == []
    assert recover_phone_values("0659327749") == []
    assert recover_whatsapp_value("+966563846414") == "https://wa.me/966563846414"
    assert recover_whatsapp_value("wa.me/966544584458") == "https://wa.me/966544584458"
    assert recover_whatsapp_value("https://wa.me/9660568562000") == "https://wa.me/966568562000"
    assert recover_whatsapp_value("https://wa.link/6ff95o") == ""
    assert recover_whatsapp_value("https://iwtsp.com/966595685732") == ""
    assert recover_linkedin_value(
        "https://www.linkedin.com/company/hawaz/about/"
    ) == "https://www.linkedin.com/company/hawaz"
    assert recover_linkedin_value("company/t7com") == "https://www.linkedin.com/company/t7com"
    assert recover_linkedin_value("https://www.linkedin.com/company/grails-com") == ""
    assert recover_linkedin_value("https://www.linkedin.com/in/zid-med-082047307/") == ""
    assert recover_twitter_value(
        "https://x.com/infinity_sollar/status/1939293178747670737?s=46"
    ) == "https://x.com/infinity_sollar"
    assert normalize_patch_value("لينكدإن", "https://www.linkedin.com/company/grails-com") == ""
    assert recover_ready_values("واتساب", "https://www.whatsapp.com/channel/abc") == []


def test_multi_phone_cell_expands_into_additive_column():
    patches = build_account_patches(
        [_ready("MA-1", "جوال", "0112782590 / 0508040404")],
        [],
    )
    assert patches["MA-1"].values_by_column["AgentReach_Phones"] == [
        "+966508040404",
        "+966112782590",
    ]


def test_refuses_existing_additive_column(tmp_path: Path):
    path = tmp_path / "bad.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["Master Account ID", "AgentReach_Phones"],
        )
        writer.writeheader()
        writer.writerow({"Master Account ID": "MA-1", "AgentReach_Phones": "keep"})
    with pytest.raises(ValueError, match="already exist"):
        write_enriched_master(path, tmp_path / "out.csv", {})
