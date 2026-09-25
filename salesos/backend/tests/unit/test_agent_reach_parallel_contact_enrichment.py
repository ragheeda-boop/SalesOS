"""Unit tests for the parallel-safe Agent Reach contact-enrichment coordinator."""

from __future__ import annotations

import csv
import json
import sys
import tempfile
import time
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import pytest

BACKEND_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = BACKEND_ROOT / "scripts"
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from agent_reach_contact_enrichment_loop import (  # noqa: E402
    async_main,
    sanitize_existing_workbook,
    worker_writes_workbook,
)
from agent_reach_parallel_contact_enrichment import (  # noqa: E402
    WorkerResult,
    build_pending_candidates,
    collect_worker_results,
    launch_workers,
    load_checkpoint_attempted_ids,
    load_worker_result,
    merge_worker_results,
    shards_overlap,
    split_shards,
    write_merge_outputs,
    write_plan,
)
from agent_reach_parallel_contact_enrichment import (  # noqa: E402
    main as coordinator_main,
)

from app.modules.agent_reach.contact_enrichment import (  # noqa: E402
    ISSUES_SHEET,
    QUEUE_HEADERS,
    READY_SHEET,
    EnrichmentCandidate,
    EnrichmentRow,
    QualityIssue,
    load_processed_accounts,
    load_queue_candidates,
)
from app.modules.agent_reach.models import AgentReachResult, ChannelType  # noqa: E402


def _candidate(account_id: str, domain: str | None = None) -> EnrichmentCandidate:
    suffix = account_id.lower().replace("-", "")
    return EnrichmentCandidate(
        account_id=account_id,
        company_name=f"Company {account_id}",
        domain=domain or f"{suffix}.com",
        tier="TIER C",
        city="Riyadh",
        has_phone=False,
        has_social=False,
    )


def _write_master_csv(rows: list[dict], path: Path | None = None) -> Path:
    csv_path = path or (Path(tempfile.mkdtemp()) / "master.csv")
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "Master Account ID",
                "Canonical_Company_Name",
                "Primary_Domain",
                "Account_Tier_v2",
                "City",
                "has_phone",
                "has_social",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)
    return csv_path


def _write_workbook(path: Path, ready_ids: list[str], issue_ids: list[str] | None = None) -> Path:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ready = wb.create_sheet(READY_SHEET)
    ready.append(["رقم الحساب", "اسم الشركة", "النطاق", "الفئة", "الحقل", "القيمة/الرابط"])
    for account_id in ready_ids:
        ready.append(
            [
                account_id,
                f"Company {account_id}",
                f"{account_id.lower()}.com",
                "TIER C",
                "جوال",
                "+966501234567",
            ]
        )
    issues = wb.create_sheet(ISSUES_SHEET)
    issues.append(["رقم الحساب", "اسم الشركة", "النطاق المسجل", "المشكلة"])
    for account_id in issue_ids or []:
        issues.append(
            [
                account_id,
                f"Company {account_id}",
                f"{account_id.lower()}.com",
                "تعذر القراءة",
            ]
        )
    wb.save(path)
    return path


def _master_row(account_id: str, domain: str | None = None) -> dict:
    return {
        "Master Account ID": account_id,
        "Canonical_Company_Name": f"Company {account_id}",
        "Primary_Domain": domain or f"{account_id.lower()}.com",
        "Account_Tier_v2": "TIER C",
        "City": "Riyadh",
        "has_phone": "0",
        "has_social": "0",
    }


def _write_ready_csv(path: Path, rows: list[EnrichmentRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["رقم الحساب", "اسم الشركة", "النطاق", "الفئة", "الحقل", "القيمة/الرابط"])
        for row in rows:
            writer.writerow(row.as_csv_row())


def _write_issues_csv(path: Path, issues: list[QualityIssue]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["رقم الحساب", "اسم الشركة", "النطاق المسجل", "المشكلة"])
        for issue in issues:
            writer.writerow(issue.as_csv_row())


def _write_checkpoint(path: Path, account_ids: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for account_id in account_ids:
            handle.write(
                json.dumps(
                    {
                        "account_id": account_id,
                        "status": "ready",
                        "ready_rows": 1,
                    },
                    ensure_ascii=False,
                )
                + "\n"
            )


def _seed_worker(
    worker_dir: Path,
    *,
    assigned: list[str],
    attempted: list[str],
    ready: list[EnrichmentRow],
    issues: list[QualityIssue],
) -> None:
    worker_dir.mkdir(parents=True, exist_ok=True)
    with (worker_dir / "assigned_queue.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(QUEUE_HEADERS)
        for account_id in assigned:
            writer.writerow(
                [
                    account_id,
                    f"Company {account_id}",
                    f"{account_id.lower()}.com",
                    "TIER C",
                    "Riyadh",
                    "False",
                    "False",
                ]
            )
    _write_checkpoint(worker_dir / "agent_reach_checkpoint.jsonl", attempted)
    _write_ready_csv(worker_dir / "agent_reach_ready_rows.csv", ready)
    _write_issues_csv(worker_dir / "agent_reach_quality_issues.csv", issues)


# ── sharding ──────────────────────────────────────────────────────


def test_split_shards_do_not_overlap():
    workers = 4
    shard_size = 25
    total = workers * shard_size
    candidates = [_candidate(f"A{index:03d}") for index in range(1, total + 1)]
    shards = split_shards(candidates, workers=workers, shard_size=shard_size)
    assert len(shards) == workers
    assert all(len(shard) == shard_size for shard in shards)
    assert shards_overlap(shards) is False
    assigned = [item.account_id for shard in shards for item in shard]
    assert assigned == [item.account_id for item in candidates[:total]]
    assert len(set(assigned)) == total


def test_split_shards_partial_last_worker():
    candidates = [_candidate(f"A{index}") for index in range(1, 6)]
    shards = split_shards(candidates, workers=3, shard_size=2)
    assert [len(shard) for shard in shards] == [2, 2, 1]
    assert shards_overlap(shards) is False


def test_split_shards_eight_workers_cover_remainder_exactly():
    workers = 8
    shard_size = 25
    remaining = 99
    candidates = [_candidate(f"A{index:03d}") for index in range(1, remaining + 1)]
    shards = split_shards(candidates, workers=workers, shard_size=shard_size)
    assigned = [item.account_id for shard in shards for item in shard]
    assert shards_overlap(shards) is False
    assert assigned == [item.account_id for item in candidates]
    assert len(assigned) == remaining
    assert sum(len(shard) for shard in shards) == remaining


def test_launch_workers_staggers_subprocess_starts(monkeypatch):
    launches: list[float] = []

    class _FakeProc:
        def poll(self) -> int:
            return 0

        def wait(self) -> int:
            return 0

    def _fake_popen(command: list[str]) -> _FakeProc:
        _ = command
        launches.append(time.monotonic())
        return _FakeProc()

    monkeypatch.setattr(
        "agent_reach_parallel_contact_enrichment.subprocess.Popen",
        _fake_popen,
    )
    stagger = 0.05
    codes = launch_workers(
        [["one"], ["two"], ["three"]],
        max_workers=8,
        stagger_seconds=stagger,
    )
    assert codes == [0, 0, 0]
    expected_launches = 3
    assert len(launches) == expected_launches
    assert launches[1] - launches[0] >= stagger * 0.8
    assert launches[2] - launches[1] >= stagger * 0.8


def test_worker_writes_workbook_false_for_queue_csv():
    assert worker_writes_workbook(queue_csv=Path("shard_001.csv"), no_workbook=False) is False
    assert worker_writes_workbook(queue_csv=Path("shard_001.csv"), no_workbook=True) is False
    assert worker_writes_workbook(queue_csv=None, no_workbook=True) is False
    assert worker_writes_workbook(queue_csv=None, no_workbook=False) is True


def test_load_queue_candidates_preserves_order():
    tmpdir = Path(tempfile.mkdtemp())
    path = tmpdir / "shard_001.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(QUEUE_HEADERS)
        writer.writerow(["A2", "Beta", "beta.com", "TIER C", "Jeddah", "True", "False"])
        writer.writerow(["A1", "Acme", "acme.com", "TIER C", "Riyadh", "False", "False"])
    loaded = load_queue_candidates(path)
    assert [item.account_id for item in loaded] == ["A2", "A1"]
    assert loaded[0].has_phone is True


# ── worker mode does not write workbook ───────────────────────────


@pytest.mark.asyncio
async def test_worker_queue_csv_does_not_write_workbook(monkeypatch):
    tmpdir = Path(tempfile.mkdtemp())
    queue = tmpdir / "shard_001.csv"
    existing = tmpdir / "existing.xlsx"
    output_dir = tmpdir / "worker_001"
    _write_workbook(existing, ["KEEP"])
    with queue.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(QUEUE_HEADERS)
        writer.writerow(["A2", "Beta", "beta.com", "TIER C", "", "False", "False"])

    page = (
        "Welcome to Beta. Call +966501234567 for help. "
        "Follow https://facebook.com/beta and keep reading this page. "
    )
    mock_service = MagicMock()
    mock_service.read_web_page = AsyncMock(
        return_value=AgentReachResult(
            success=True,
            channel=ChannelType.WEB,
            action="read",
            data=page,
        )
    )

    import agent_reach_contact_enrichment_loop as loop_mod

    def _service_factory(timeout_seconds: int = 30) -> MagicMock:
        _ = timeout_seconds
        return mock_service

    monkeypatch.setattr(loop_mod, "AgentReachService", _service_factory)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "agent_reach_contact_enrichment_loop.py",
            "--queue-csv",
            str(queue),
            "--existing-xlsx",
            str(existing),
            "--output-dir",
            str(output_dir),
            "--sleep-seconds",
            "0",
        ],
    )

    assert await async_main() == 0
    assert (output_dir / "agent_reach_ready_rows.csv").exists()
    assert (output_dir / "agent_reach_quality_issues.csv").exists()
    assert (output_dir / "agent_reach_checkpoint.jsonl").exists()
    assert (output_dir / "summary.json").exists()
    assert list(output_dir.glob("*.xlsx")) == []
    summary = json.loads((output_dir / "summary.json").read_text(encoding="utf-8"))
    assert summary["workbook"] is None
    assert summary["mode"] == "worker"
    assert summary["incomplete"] is False


# ── coordinator merge ─────────────────────────────────────────────


def _ready(account_id: str, value: str = "+966501234567") -> EnrichmentRow:
    return EnrichmentRow(
        account_id,
        f"Company {account_id}",
        f"{account_id.lower()}.com",
        "TIER C",
        "جوال",
        value,
    )


def _issue(account_id: str, message: str = "تعذر القراءة") -> QualityIssue:
    return QualityIssue(account_id, f"Company {account_id}", f"{account_id.lower()}.com", message)


def test_merge_appends_each_ready_and_issue_row_once():
    tmpdir = Path(tempfile.mkdtemp())
    existing = tmpdir / "existing.xlsx"
    output_dir = tmpdir / "parallel"
    _write_workbook(existing, ["A1"])

    worker_1 = output_dir / "workers" / "worker_001"
    worker_2 = output_dir / "workers" / "worker_002"
    shared = _ready("A2")
    _seed_worker(
        worker_1,
        assigned=["A2", "A3"],
        attempted=["A2", "A3"],
        ready=[shared],
        issues=[_issue("A3")],
    )
    _seed_worker(
        worker_2,
        assigned=["A4", "A5"],
        attempted=["A4"],
        ready=[shared, _ready("A4")],
        issues=[],
    )

    bundle = merge_worker_results(collect_worker_results(output_dir))
    assert [row.account_id for row in bundle.ready_rows] == ["A2", "A4"]
    assert [issue.account_id for issue in bundle.issues] == ["A3"]
    assert bundle.incomplete_shards == ["worker_002"]

    summary = write_merge_outputs(output_dir, bundle, existing_xlsx=existing, pending_total=10)
    workbook = Path(summary["workbook"])
    processed = load_processed_accounts(workbook)
    assert processed == {"A1", "A2", "A3", "A4"}
    assert "A5" not in processed
    expected_ready = 2
    expected_issues = 1
    assert summary["ready_rows"] == expected_ready
    assert summary["issue_rows"] == expected_issues
    assert summary["next_existing_xlsx"] == str(workbook.resolve())


def test_incomplete_shards_reported_from_checkpoint_and_assigned_queue():
    tmpdir = Path(tempfile.mkdtemp())
    worker_dir = tmpdir / "worker_003"
    _seed_worker(
        worker_dir,
        assigned=["A8", "A9"],
        attempted=["A8"],
        ready=[_ready("A8")],
        issues=[],
    )
    result = load_worker_result(worker_dir, shard_id="shard_003")
    assert result.incomplete is True
    assert result.unattempted_ids == ["A9"]
    assert result.attempted_ids == ["A8"]

    bundle = merge_worker_results([result])
    assert bundle.incomplete_shards == ["shard_003"]
    assert bundle.attempted_ids == ["A8"]
    assert "A9" not in bundle.attempted_ids


def test_checkpoint_only_attempt_is_processed_unattempted_is_not():
    tmpdir = Path(tempfile.mkdtemp())
    path = tmpdir / "agent_reach_checkpoint.jsonl"
    _write_checkpoint(path, ["A8"])
    assert load_checkpoint_attempted_ids(path) == ["A8"]

    worker_dir = tmpdir / "worker_empty_csv"
    worker_dir.mkdir()
    _write_checkpoint(worker_dir / "agent_reach_checkpoint.jsonl", ["A8"])
    with (worker_dir / "assigned_queue.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(QUEUE_HEADERS)
        writer.writerow(["A8", "Co", "a8.com", "TIER C", "", "False", "False"])
        writer.writerow(["A9", "Co", "a9.com", "TIER C", "", "False", "False"])
    result = load_worker_result(worker_dir, shard_id="shard_empty")
    assert result.attempted_ids == ["A8"]
    assert result.unattempted_ids == ["A9"]
    assert result.incomplete is True


def test_merged_workbook_becomes_next_resume_input():
    tmpdir = Path(tempfile.mkdtemp())
    master = _write_master_csv(
        [
            _master_row("A1", "a1.com"),
            _master_row("A2", "a2.com"),
            _master_row("A3", "a3.com"),
            _master_row("A4", "a4.com"),
            _master_row("A5", "a5.com"),
            _master_row("A6", "a6.com"),
        ]
    )
    existing = _write_workbook(tmpdir / "existing.xlsx", ["A1"])
    output_dir = tmpdir / "parallel"
    _seed_worker(
        output_dir / "workers" / "worker_001",
        assigned=["A2", "A3"],
        attempted=["A2", "A3"],
        ready=[_ready("A2")],
        issues=[_issue("A3")],
    )
    _seed_worker(
        output_dir / "workers" / "worker_002",
        assigned=["A4", "A5"],
        attempted=["A4"],
        ready=[_ready("A4")],
        issues=[],
    )

    bundle = merge_worker_results(collect_worker_results(output_dir))
    summary = write_merge_outputs(output_dir, bundle, existing_xlsx=existing, pending_total=5)
    next_xlsx = Path(summary["next_existing_xlsx"])

    next_pending, next_processed = build_pending_candidates(
        master,
        next_xlsx,
        tiers={"TIER C"},
    )
    next_ids = [item.account_id for item in next_pending]
    assert next_processed == {"A1", "A2", "A3", "A4"}
    assert "A2" not in next_ids
    assert "A3" not in next_ids
    assert "A4" not in next_ids
    assert "A5" in next_ids
    assert "A6" in next_ids


def test_write_plan_creates_non_overlapping_shards_and_commands():
    tmpdir = Path(tempfile.mkdtemp())
    master = _write_master_csv(
        [_master_row(f"A{index}", f"co{index}.com") for index in range(1, 8)]
    )
    existing = _write_workbook(tmpdir / "existing.xlsx", ["A1"])
    pending, processed = build_pending_candidates(master, existing, tiers={"TIER C"})
    output_dir = tmpdir / "plan"
    plan = write_plan(
        output_dir,
        existing_xlsx=existing,
        master_csv=master,
        pending=pending,
        processed=processed,
        workers=2,
        shard_size=2,
        sleep_seconds=0,
    )
    expected_assigned = 4
    expected_pending = 6
    expected_commands = 2
    assert plan["assigned_this_run"] == expected_assigned
    assert plan["pending_queue_total"] == expected_pending
    assert len(plan["worker_commands"]) == expected_commands
    shard_ids = []
    for record in plan["shards"]:
        loaded = load_queue_candidates(Path(record["shard_path"]))
        shard_ids.append({item.account_id for item in loaded})
        assert "--no-workbook" in record["command"]
        assert "--queue-csv" in record["command"]
        assert Path(record["worker_dir"], "assigned_queue.csv").exists()
    assert shard_ids[0].isdisjoint(shard_ids[1])
    assert (output_dir / "pending_queue.csv").exists()
    assert (output_dir / "worker_commands.txt").exists()


def test_unattempted_rows_are_not_inferred_as_processed():
    result = WorkerResult(
        shard_id="shard_001",
        worker_dir=Path("unused"),
        assigned_ids=["A1", "A2"],
        attempted_ids=["A1"],
        ready_rows=[_ready("A1")],
        incomplete=True,
    )
    bundle = merge_worker_results([result])
    assert bundle.attempted_ids == ["A1"]
    assert "A2" not in {row.account_id for row in bundle.ready_rows}
    assert "A2" not in {issue.account_id for issue in bundle.issues}


def test_sanitize_existing_workbook_copies_and_cleans_ready_urls():
    from openpyxl import load_workbook

    tmpdir = Path(tempfile.mkdtemp())
    source = tmpdir / "source.xlsx"
    dest = tmpdir / "cleaned" / "cleaned.xlsx"
    _write_workbook(source, ["KEEP"], issue_ids=["ISSUE"])
    wb = load_workbook(source)
    ready = wb[READY_SHEET]
    ready.append(
        [
            "MA-0047641",
            "Hospital",
            "example.com",
            "TIER C",
            "واتساب",
            "https://api.whatsapp.com/send?phone=966920012777&text=)![Image",
        ]
    )
    ready.append(
        [
            "MA-0051488",
            "Contractor",
            "example.org",
            "TIER C",
            "واتساب",
            "https://wa.me/966566222773)![Image",
        ]
    )
    ready.append(
        [
            "MA-0237577",
            "Madaf",
            "madaf.com",
            "TIER C",
            "لينكدإن",
            "https://www.linkedin.com/company/madaf-trading-and-contracting-co.-ltd.",
        ]
    )
    wb[ISSUES_SHEET].append(
        ["MA-0099999", "Other", "other.com", "text mentions ![Image but is not a URL"]
    )
    wb.save(source)
    source_mtime = source.stat().st_mtime
    source_size = source.stat().st_size

    stats = sanitize_existing_workbook(source, dest)
    expected_ready_rows = 4
    expected_issue_rows = 2
    expected_cleaned = 2
    assert dest.exists()
    assert source.stat().st_mtime == source_mtime
    assert source.stat().st_size == source_size
    assert stats["ready_rows"] == expected_ready_rows
    assert stats["issue_rows"] == expected_issue_rows
    assert stats["image_artifacts_ready_before"] == expected_cleaned
    assert stats["image_artifacts_ready_after"] == 0
    assert stats["ready_urls_cleaned"] == expected_cleaned

    cleaned = load_workbook(dest, data_only=True)
    values = [row[5] for row in cleaned[READY_SHEET].iter_rows(min_row=2, values_only=True)]
    assert "https://api.whatsapp.com/send?phone=966920012777&text=" in values
    assert "https://wa.me/966566222773" in values
    assert (
        "https://www.linkedin.com/company/madaf-trading-and-contracting-co.-ltd."
        in values
    )
    assert all("![Image" not in str(value or "") for value in values)
    issue_text = [
        row[3] for row in cleaned[ISSUES_SHEET].iter_rows(min_row=2, values_only=True)
    ]
    assert "text mentions ![Image but is not a URL" in issue_text
    cleaned.close()


def test_sanitize_existing_workbook_refuses_in_place_overwrite():
    tmpdir = Path(tempfile.mkdtemp())
    source = tmpdir / "same.xlsx"
    _write_workbook(source, ["KEEP"])
    with pytest.raises(ValueError, match="in-place"):
        sanitize_existing_workbook(source, source)


def test_coordinator_sanitize_xlsx_writes_new_workbook():
    tmpdir = Path(tempfile.mkdtemp())
    source = tmpdir / "source.xlsx"
    output_dir = tmpdir / "cleaned"
    _write_workbook(source, ["KEEP"])
    from openpyxl import load_workbook

    wb = load_workbook(source)
    wb[READY_SHEET].append(
        [
            "MA-0051488",
            "Contractor",
            "example.org",
            "TIER C",
            "واتساب",
            "https://wa.me/966566222773)![Image",
        ]
    )
    wb.save(source)

    assert coordinator_main(
        [
            "--sanitize-xlsx",
            "--existing-xlsx",
            str(source),
            "--output-dir",
            str(output_dir),
        ]
    ) == 0
    summary_path = output_dir / "sanitize_summary.json"
    assert summary_path.exists()
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    assert summary["mode"] == "sanitize"
    assert summary["image_artifacts_ready_after"] == 0
    assert (output_dir / "14_Website_Phone_Social_Enrichment_AgentReach.xlsx").exists()


def test_coordinator_rejects_sanitize_with_plan_only():
    tmpdir = Path(tempfile.mkdtemp())
    usage_error = 2
    assert (
        coordinator_main(
            [
                "--sanitize-xlsx",
                "--plan-only",
                "--existing-xlsx",
                str(tmpdir / "missing.xlsx"),
                "--output-dir",
                str(tmpdir / "out"),
            ]
        )
        == usage_error
    )
