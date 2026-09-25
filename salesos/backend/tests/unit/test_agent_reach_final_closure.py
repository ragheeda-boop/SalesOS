"""Unit tests for final Agent Reach closure helpers."""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook

BACKEND_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = BACKEND_ROOT / "scripts"
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from agent_reach_final_closure import (  # noqa: E402
    METHOD_DOMAIN,
    METHOD_EXTERNAL,
    METHOD_HUMAN,
    METHOD_INSUFFICIENT,
    classify_next_method,
    clean_ready_workbook,
)
from agent_reach_master_enrichment_patch import ReadyRecord  # noqa: E402

from app.modules.agent_reach.contact_enrichment import ISSUES_SHEET, READY_SHEET  # noqa: E402


def _write_workbook(path: Path, ready: list[ReadyRecord]) -> Path:
    wb = Workbook()
    ready_ws = wb.create_sheet(READY_SHEET)
    ready_ws.append(["رقم الحساب", "اسم الشركة", "النطاق", "الفئة", "الحقل", "القيمة/الرابط"])
    for row in ready:
        ready_ws.append(
            [row.account_id, row.company_name, row.domain, row.tier, row.field, row.value]
        )
    issue_ws = wb.create_sheet(ISSUES_SHEET)
    issue_ws.append(["رقم الحساب", "اسم الشركة", "النطاق المسجل", "المشكلة"])
    default = wb[wb.sheetnames[0]]
    if default.title not in {READY_SHEET, ISSUES_SHEET}:
        wb.remove(default)
    wb.save(path)
    return path


def test_classify_next_method_priority():
    assert classify_next_method(
        categories={"typo_domain"},
        status="",
        attempted=True,
        usable_domain=False,
        master_has_phone=False,
        master_has_social=False,
        missing_any=True,
    ) == METHOD_DOMAIN
    assert classify_next_method(
        categories={"company_domain_mismatch"},
        status="PENDING_REVIEW",
        attempted=True,
        usable_domain=True,
        master_has_phone=False,
        master_has_social=False,
        missing_any=True,
    ) == METHOD_HUMAN
    assert classify_next_method(
        categories={"no_contact_found"},
        status="QUALITY_BACKLOG",
        attempted=True,
        usable_domain=True,
        master_has_phone=False,
        master_has_social=False,
        missing_any=True,
    ) == METHOD_EXTERNAL
    assert classify_next_method(
        categories=set(),
        status="",
        attempted=False,
        usable_domain=False,
        master_has_phone=False,
        master_has_social=False,
        missing_any=True,
    ) == METHOD_INSUFFICIENT


def test_clean_ready_workbook_recovers_and_deletes(tmp_path: Path):
    source = _write_workbook(
        tmp_path / "src.xlsx",
        [
            ReadyRecord("MA-1", "A", "a.com", "TIER C", "جوال", "+966501234567"),
            ReadyRecord("MA-2", "B", "b.com", "TIER C", "واتساب", "https://wa.link/abc"),
            ReadyRecord("MA-3", "C", "c.com", "TIER C", "جوال", "0501112233 / 0552223344"),
            ReadyRecord(
                "MA-4",
                "D",
                "d.com",
                "ANTI-ICP",
                "لينكدإن",
                "https://www.linkedin.com/company/acme/about/",
            ),
        ],
    )
    dest = tmp_path / "cleaned.xlsx"
    stats = clean_ready_workbook(source, dest)
    assert dest.exists()
    assert source.exists()
    assert stats["deleted_rows"] == 1
    expected_ready_before = 4
    expected_ready_after = 4  # 1 deleted, 1 split extra
    assert stats["ready_rows_before"] == expected_ready_before
    assert stats["ready_rows_after"] == expected_ready_after
    from openpyxl import load_workbook

    wb = load_workbook(dest, read_only=True, data_only=True)
    rows = list(wb[READY_SHEET].iter_rows(min_row=2, values_only=True))
    wb.close()
    values = {(str(row[0]), str(row[4]), str(row[5])) for row in rows if row and row[0]}
    assert ("MA-2", "واتساب", "https://wa.link/abc") not in values
    assert ("MA-1", "جوال", "+966501234567") in values
    assert ("MA-3", "جوال", "+966501112233") in values
    assert ("MA-3", "جوال", "+966552223344") in values
    assert ("MA-4", "لينكدإن", "https://www.linkedin.com/company/acme") in values


def test_backlog_csv_writer_roundtrip(tmp_path: Path):
    from agent_reach_final_closure import BACKLOG_HEADERS, _write_csv

    path = tmp_path / "backlog.csv"
    _write_csv(
        path,
        BACKLOG_HEADERS,
        [
            {
                "Master Account ID": "MA-1",
                "company_name": "Acme",
                "tier": "TIER C",
                "tier_bucket": "TIER C",
                "primary_domain": "acme.com",
                "usable_domain": "true",
                "attempted_agent_reach": "true",
                "enrichment_status": "PENDING_REVIEW",
                "missing_phone": "false",
                "missing_whatsapp": "true",
                "missing_linkedin": "true",
                "missing_other_social": "true",
                "invalid_dead_or_mismatch_site": "false",
                "quality_issues_only": "false",
                "no_recoverable_website_social_evidence": "false",
                "quality_categories": "",
                "next_method": METHOD_EXTERNAL,
                "notes": "agent_reach_already_attempted",
            }
        ],
    )
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0]["Master Account ID"] == "MA-1"
    assert rows[0]["missing_whatsapp"] == "true"
