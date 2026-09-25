"""Website contact enrichment helpers for Agent Reach.

This module keeps the MUHIDE spreadsheet continuation as a review-only flow:
read public company websites through Agent Reach, extract contact/social fields,
and emit candidate rows for human review before any master-data write.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import parse_qs, urlparse, urlunparse

from .models import AgentReachResult
from .service import AgentReachService

READY_SHEET = "مرشحون جاهزون"
ISSUES_SHEET = "مشاكل جودة بيانات مكتشفة"
QUEUE_HEADERS = [
    "account_id",
    "company_name",
    "domain",
    "tier",
    "city",
    "has_phone",
    "has_social",
]

SOCIAL_FIELDS: tuple[tuple[str, str], ...] = (
    ("فيسبوك", "facebook.com"),
    ("انستقرام", "instagram.com"),
    ("تويتر/X", "x.com"),
    ("تويتر/X", "twitter.com"),
    ("لينكدإن", "linkedin.com"),
    ("واتساب", "wa.me"),
    ("واتساب", "whatsapp.com"),
    ("تيك توك", "tiktok.com"),
    ("يوتيوب", "youtube.com"),
    ("سناب شات", "snapchat.com"),
)

GENERIC_DOMAINS = {
    "facebook.com",
    "gmail.com",
    "hotmail.com",
    "icloud.com",
    "instagram.com",
    "linkedin.com",
    "linktr.ee",
    "live.com",
    "outlook.com",
    "twitter.com",
    "wa.me",
    "whatsapp.com",
    "x.com",
    "yahoo.com",
    "ymail.com",
    "youtube.com",
}

GENERIC_SUFFIXES = (
    ".facebook.com",
    ".gmail.com",
    ".hotmail.com",
    ".icloud.com",
    ".instagram.com",
    ".linkedin.com",
    ".live.com",
    ".outlook.com",
    ".twitter.com",
    ".yahoo.com",
    ".ymail.com",
    ".youtube.com",
)

MARKDOWN_IMAGE_ARTIFACT = "![Image"

SHARED_DOMAIN_THRESHOLD = 3
MIN_PAGE_TEXT_CHARS = 80
CONCATENATED_HTTP_MIN = 2
LINKEDIN_MIN_PATH_PARTS = 2
LINKEDIN_MAX_PATH_PARTS = 2
TRAILING_URL_GARBAGE = r"[).,*؛،\"'\u00b7\u0621-\u064a\u0660-\u0669\[\]!#]+"
SA_INTL_LENGTHS = {11, 12}
SA_MOBILE_LOCAL_LENGTH = 10
SA_MOBILE_SHORT_LENGTH = 9
SA_TOLL_FREE_LENGTHS = {9, 10}
SA_LANDLINE_LENGTH = 10

TRACKING_SUBDOMAINS = frozenset({
    "analytics",
    "pixel",
    "ads",
    "advertising",
    "connect",
    "developers",
    "help",
    "mail",
    "media",
    "static",
    "upload",
})

TWITTER_RESERVED_FIRST_SEGMENTS = frozenset({
    "compose",
    "explore",
    "hashtag",
    "home",
    "i",
    "intent",
    "login",
    "search",
    "share",
    "signup",
})

LINKEDIN_ALLOWED_FIRST_SEGMENTS = frozenset({
    "company",
    "in",
    "school",
    "showcase",
})

# Website-platform / CMS vendor pages that leak from theme footers.
# Match the first hyphen/underscore token of the LinkedIn entity slug.
PLATFORM_LINKEDIN_SLUGS = frozenset({
    "absher",
    "acmethemes",
    "bigcommerce",
    "drupal",
    "elementor",
    "expandcart",
    "godaddy",
    "grails",
    "hostinger",
    "joomla",
    "magento",
    "networksolutions",
    "odoo",
    "prestashop",
    "protonprivacy",
    "salla",
    "shopify",
    "squarespace",
    "webflow",
    "weebly",
    "wix",
    "woocommerce",
    "wordpress",
    "zid",
})

INSTAGRAM_RESERVED_FIRST_SEGMENTS = frozenset({
    "accounts",
    "explore",
    "p",
    "reel",
    "reels",
    "stories",
    "tv",
})


def _domain_slug(domain: str) -> str:
    """Extract a normalised slug from a company domain for plausibility checks."""
    slug = domain.split(".")[0]
    slug = re.sub(r"[^a-z0-9]", "", slug.lower())
    return slug


def _safe_urlparse(url: str):
    """Parse a URL, or return None for malformed values such as unclosed IPv6."""
    try:
        return urlparse(url)
    except ValueError:
        return None


def _social_username(url: str) -> str:
    """Extract the username/handle from a social-media URL path."""
    parsed = _safe_urlparse(url)
    if parsed is None:
        return ""
    path = (parsed.path or "").strip("/")
    if not path:
        return ""
    username = path.split("/")[0]
    username = re.sub(r"[^a-z0-9._]", "", username.lower())
    return username


def _is_social_plausible(username: str, company_slug: str) -> bool:
    """Return True if a social handle could plausibly belong to the company."""
    if not username or not company_slug:
        return False
    if username == company_slug:
        return True
    min_token = 3
    if len(username) < min_token or len(company_slug) < min_token:
        return False
    if company_slug in username or username in company_slug:
        return True
    return False


@dataclass(frozen=True)
class EnrichmentCandidate:
    account_id: str
    company_name: str
    domain: str
    tier: str
    city: str = ""
    has_phone: bool = False
    has_social: bool = False

    @property
    def tier_bucket(self) -> str:
        return tier_bucket(self.tier)


@dataclass(frozen=True)
class EnrichmentRow:
    account_id: str
    company_name: str
    domain: str
    tier: str
    field: str
    value: str

    def as_csv_row(self) -> list[str]:
        return [
            self.account_id,
            self.company_name,
            self.domain,
            self.tier,
            self.field,
            self.value,
        ]


@dataclass(frozen=True)
class QualityIssue:
    account_id: str
    company_name: str
    domain: str
    issue: str

    def as_csv_row(self) -> list[str]:
        return [self.account_id, self.company_name, self.domain, self.issue]


def normalize_domain(value: str | None) -> str:
    """Normalize a website domain without treating it as identity proof."""
    if not value:
        return ""
    domain = str(value).strip().lower()
    domain = re.sub(r"^https?://", "", domain)
    domain = re.sub(r"^www\.", "", domain)
    domain = domain.split("/", 1)[0].strip()
    return domain


def is_safe_company_domain(domain: str, domain_counts: Counter[str] | None = None) -> bool:
    """Return whether a domain is suitable for one-company website enrichment."""
    domain = normalize_domain(domain)
    if not domain or "." not in domain:
        return False
    if domain in GENERIC_DOMAINS or domain.endswith(GENERIC_SUFFIXES):
        return False
    if domain_counts is not None and domain_counts[domain] >= SHARED_DOMAIN_THRESHOLD:
        return False
    return True


def tier_bucket(value: str | None) -> str:
    raw = (value or "").strip()
    upper = raw.upper()
    if upper == "CLASS A":
        return "CLASS A"
    if upper == "TIER B":
        return "TIER B"
    if upper.startswith("TIER C"):
        return "TIER C"
    if upper.startswith("TIER D"):
        return "TIER D"
    if upper == "ANTI-ICP":
        return "ANTI-ICP"
    return raw or "UNKNOWN"


def parse_bool(value: object) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "y"}


def load_processed_accounts(workbook_path: Path) -> set[str]:
    """Load account IDs already present in the in-progress workbook."""
    from openpyxl import load_workbook

    processed: set[str] = set()
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    for sheet_name in (READY_SHEET, ISSUES_SHEET):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            account_id = str(row[0] or "").strip()
            if account_id:
                processed.add(account_id)
    return processed


def load_master_candidates(
    master_csv_path: Path,
    *,
    processed_accounts: set[str] | None = None,
    tiers: set[str] | None = None,
) -> list[EnrichmentCandidate]:
    """Build the safe, unprocessed enrichment queue from MUHIDE master CSV."""
    processed_accounts = processed_accounts or set()
    rows: list[dict[str, str]] = []
    domain_counts: Counter[str] = Counter()

    with master_csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            domain = normalize_domain(row.get("Primary_Domain"))
            row["_domain"] = domain
            rows.append(row)
            if domain:
                domain_counts[domain] += 1

    candidates: list[EnrichmentCandidate] = []
    for row in rows:
        account_id = (row.get("Master Account ID") or "").strip()
        if not account_id or account_id in processed_accounts:
            continue

        domain = row.get("_domain") or ""
        if not is_safe_company_domain(domain, domain_counts):
            continue

        has_phone = parse_bool(row.get("has_phone"))
        has_social = parse_bool(row.get("has_social"))
        if has_phone and has_social:
            continue

        bucket = tier_bucket(row.get("Account_Tier_v2"))
        if tiers and bucket not in tiers:
            continue

        candidates.append(
            EnrichmentCandidate(
                account_id=account_id,
                company_name=(row.get("Canonical_Company_Name") or "").strip(),
                domain=domain,
                tier=row.get("Account_Tier_v2") or bucket,
                city=(row.get("City") or "").strip(),
                has_phone=has_phone,
                has_social=has_social,
            )
        )

    return sorted(candidates, key=lambda item: item.account_id)


def queue_row_from_candidate(candidate: EnrichmentCandidate) -> list[str]:
    """Serialize one queue/shard row in the worker CSV contract."""
    return [
        candidate.account_id,
        candidate.company_name,
        candidate.domain,
        candidate.tier_bucket,
        candidate.city,
        str(candidate.has_phone),
        str(candidate.has_social),
    ]


def load_queue_candidates(queue_csv_path: Path) -> list[EnrichmentCandidate]:
    """Load a fixed, coordinator-assigned account list for one worker.

    The shard is the source of truth. Workers do not re-filter against the
    workbook or master CSV — that would risk skipping or overlapping work.
    """
    candidates: list[EnrichmentCandidate] = []
    with queue_csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            account_id = (row.get("account_id") or "").strip()
            if not account_id:
                continue
            candidates.append(
                EnrichmentCandidate(
                    account_id=account_id,
                    company_name=(row.get("company_name") or "").strip(),
                    domain=normalize_domain(row.get("domain")),
                    tier=(row.get("tier") or "").strip(),
                    city=(row.get("city") or "").strip(),
                    has_phone=parse_bool(row.get("has_phone")),
                    has_social=parse_bool(row.get("has_social")),
                )
            )
    return candidates


def _has_concatenated_http(url: str) -> bool:
    """True when a value embeds a second http(s) URL (concatenated leak)."""
    return len(re.findall(r"https?://", url or "", flags=re.I)) >= CONCATENATED_HTTP_MIN


def _clean_url(raw_url: str) -> str:
    url = raw_url.strip()
    url = re.sub(r"!\[.*$", "", url)
    url = re.sub(TRAILING_URL_GARBAGE + r"$", "", url)
    url = url.replace("\\/", "/")
    url = re.sub(r"\)[\[].*$", "", url)
    url = re.sub(r"[).,*؛،\"'\u00b7\u0621-\u064a\u0660-\u0669\[\]!#,\s]+$", "", url)
    return url


def sanitize_ready_url_value(value: str) -> str:
    """Apply ``_clean_url`` only to HTTP(S) ready values that still leak ``![Image``.

    A blanket pass is unsafe: genuine LinkedIn slugs can end in a period or
    Arabic letters that ``_clean_url`` would strip.
    """
    raw = value if isinstance(value, str) else ("" if value is None else str(value))
    stripped = raw.strip()
    if not stripped.lower().startswith(("http://", "https://")):
        return raw
    if MARKDOWN_IMAGE_ARTIFACT not in stripped:
        return raw
    cleaned = _clean_url(stripped)
    if cleaned.lower().startswith(("http://", "https://")):
        return cleaned
    return raw


def _dedupe_urls(urls: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    unique: list[str] = []
    for raw in urls:
        url = _clean_url(raw)
        if not url.lower().startswith(("http://", "https://")):
            continue
        key = url.lower().rstrip("/")
        if key in seen:
            continue
        seen.add(key)
        unique.append(url)
    return unique


def _normalize_twitter_profile_url(url: str) -> str | None:
    """Return a clean X/Twitter profile URL, or None for share/intent/reserved paths.

    Genuine ``x.com/Handle`` / ``twitter.com/Handle`` paths are kept. Tracking
    query (``?t=``, ``&s=``) is stripped. ``/intent/`` and ``/share`` are rejected.
    """
    parsed = _safe_urlparse(url)
    if parsed is None:
        return None
    path = (parsed.path or "").strip("/")
    if not path:
        return None
    first = path.split("/")[0].lower()
    if first in TWITTER_RESERVED_FIRST_SEGMENTS:
        return None
    lowered_path = f"/{path.lower()}/"
    if "/status/" in lowered_path or "/statuses/" in lowered_path:
        return None
    return urlunparse((parsed.scheme, parsed.netloc, f"/{path.split('/')[0]}", "", "", ""))


def _linkedin_entity_token(parts: list[str]) -> str:
    """First hyphen/underscore token of a LinkedIn company/in/school/showcase slug."""
    if len(parts) < LINKEDIN_MIN_PATH_PARTS:
        return ""
    token = re.split(r"[-_/]+", parts[1].lower())[0]
    return re.sub(r"[^a-z0-9]", "", token)


def _linkedin_compact_slug(parts: list[str]) -> str:
    """Full LinkedIn slug with hyphens/underscores stripped (network-solutions)."""
    if len(parts) < LINKEDIN_MIN_PATH_PARTS:
        return ""
    return re.sub(r"[^a-z0-9]", "", parts[1].lower())


def _is_platform_linkedin_slug(parts: list[str]) -> bool:
    return (
        _linkedin_entity_token(parts) in PLATFORM_LINKEDIN_SLUGS
        or _linkedin_compact_slug(parts) in PLATFORM_LINKEDIN_SLUGS
    )


def _normalize_linkedin_url(url: str) -> str | None:
    """Keep company/profile/school/showcase LinkedIn URLs; drop email-paths and junk.

    Rejects any path containing ``@`` (mailto-style leaks such as
    ``/Rak.hr@rak-j.com``) and any first segment outside
    ``company`` / ``in`` / ``school`` / ``showcase``. Scheme-less
    ``linkedin.com/company/...`` values are treated as HTTPS.
    Platform-vendor slugs (``odoo``, ``shopify``, ``salla``,
    ``network-solutions``, …) are rejected so theme-footer links are not
    treated as the company's page.
    Extra path segments such as ``/mycompany/verification/`` are rejected.
    """
    raw = (url or "").strip()
    if not raw:
        return None
    if "://" not in raw:
        raw = f"https://{raw.lstrip('/')}"
    parsed = _safe_urlparse(raw)
    path = (parsed.path or "").strip("/") if parsed is not None else ""
    parts = path.split("/") if path else []
    first = parts[0].lower() if parts else ""
    nested_slug = any(
        part.lower().startswith("http") or "://" in part or "linkedin.com" in part.lower()
        for part in parts[1:]
    )
    if (
        parsed is None
        or not path
        or "@" in path
        or _has_concatenated_http(raw)
        or first not in LINKEDIN_ALLOWED_FIRST_SEGMENTS
        or nested_slug
        or _is_platform_linkedin_slug(parts)
        or len(parts) > LINKEDIN_MAX_PATH_PARTS
    ):
        return None
    return url


def _normalize_whatsapp_url(url: str) -> str | None:
    """Keep Saudi WhatsApp chat links; drop leftover ``text=`` / markdown query junk.

    Confirmed leak: ``&text=)Contact`` and ``&text=)![Image`` survive generic
    ``_clean_url`` because the garbage is a query value, not a trailing suffix.
    Rebuild from the validated phone only.
    """
    parsed = _safe_urlparse(url)
    if parsed is None:
        return None
    host = (parsed.netloc or "").lower()
    scheme = parsed.scheme or "https"
    if host == "wa.me" or host.endswith(".wa.me"):
        normalized = _normalize_phone(parsed.path.lstrip("/"))
        if normalized is None:
            return None
        digits = normalized.lstrip("+")
        return urlunparse((scheme, parsed.netloc, f"/{digits}", "", "", ""))
    if host == "whatsapp.com" or host.endswith(".whatsapp.com"):
        query_phone = parse_qs(parsed.query).get("phone", [""])[0]
        normalized = _normalize_phone(query_phone)
        if normalized is None:
            return None
        digits = normalized.lstrip("+")
        path = parsed.path or "/send"
        return urlunparse((scheme, parsed.netloc, path, "", f"phone={digits}", ""))
    return None


def _is_sa_whatsapp_url(url: str) -> bool:
    return _normalize_whatsapp_url(url) is not None


def _accepted_social_url(url: str, field: str, slug: str) -> str | None:
    """Return a keepable social URL, or None if the candidate should be dropped."""
    if _has_concatenated_http(url):
        return None
    if field == "واتساب":
        return _normalize_whatsapp_url(url)
    if field == "لينكدإن":
        return _normalize_linkedin_url(url)
    candidate = url
    if field == "تويتر/X":
        candidate = _normalize_twitter_profile_url(url)
        if candidate is None:
            return None
    if slug:
        username = _social_username(candidate)
        if username and not _is_social_plausible(username, slug):
            return None
    return candidate


def extract_social_links(
    text: str, *, company_domain: str = ""
) -> dict[str, list[str]]:
    """Extract supported social links from Agent Reach page text."""
    urls = _dedupe_urls(re.findall(r"https?://[^\s<>\[\]\"']+", text or ""))
    by_field: dict[str, list[str]] = {}
    slug = _domain_slug(company_domain)
    for url in urls:
        parsed = _safe_urlparse(url)
        if parsed is None:
            continue
        host = (parsed.netloc or "").lower()
        path = (parsed.path or "").rstrip("/")
        if not path:
            continue
        if "/status/" in path or "/statuses/" in path:
            continue
        if "/i/" in path or "/search" in path or "/explore" in path:
            continue
        first = path.strip("/").split("/")[0].lower()
        if "instagram.com" in host and first in INSTAGRAM_RESERVED_FIRST_SEGMENTS:
            continue
        prefix = host.split(".")[0]
        if prefix in TRACKING_SUBDOMAINS:
            continue
        for field, needle in SOCIAL_FIELDS:
            if host == needle or host.endswith(f".{needle}"):
                accepted = _accepted_social_url(url, field, slug)
                if accepted is None:
                    break
                existing = by_field.setdefault(field, [])
                key = accepted.lower().rstrip("/")
                if any(item.lower().rstrip("/") == key for item in existing):
                    break
                existing.append(accepted)
                break
    return by_field


def _normalize_phone(raw: str) -> str | None:
    digits = re.sub(r"\D", "", raw)
    if digits.startswith("00966"):
        digits = "966" + digits[5:]

    normalized = None
    if digits.startswith("966") and len(digits) in SA_INTL_LENGTHS:
        normalized = f"+{digits}"
    elif digits.startswith("05") and len(digits) == SA_MOBILE_LOCAL_LENGTH:
        normalized = "+966" + digits[1:]
    elif digits.startswith("5") and len(digits) == SA_MOBILE_SHORT_LENGTH:
        normalized = "+966" + digits
    elif (
        digits.startswith("9200")
        and len(digits) == SA_MOBILE_SHORT_LENGTH
        or digits.startswith("800")
        and len(digits) in SA_TOLL_FREE_LENGTHS
    ):
        normalized = digits
    elif (
        digits.startswith(("011", "012", "013", "014", "016", "017"))
        and len(digits) == SA_LANDLINE_LENGTH
    ):
        normalized = "+966" + digits[1:]
    return normalized


def extract_sa_phones(text: str) -> list[str]:
    """Extract plausible Saudi phone numbers from page text."""
    if not text:
        return []

    patterns = [
        r"(?:\+|00)?966[\s().-]*5[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d",
        r"\b05[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d\b",
        r"\b9200[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d\b",
        r"\b800[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d\b",
        r"\b01[1-467][\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d[\s().-]*\d\b",
    ]
    found: list[str] = []
    for pattern in patterns:
        for match in re.findall(pattern, text):
            normalized = _normalize_phone(match)
            if normalized and normalized not in found:
                found.append(normalized)
    return found


def build_enrichment_rows(candidate: EnrichmentCandidate, page_text: str) -> list[EnrichmentRow]:
    """Convert extracted page content into review-table rows."""
    rows: list[EnrichmentRow] = []
    if not candidate.has_phone:
        rows.extend(
            [
                EnrichmentRow(
                    candidate.account_id,
                    candidate.company_name,
                    candidate.domain,
                    candidate.tier_bucket,
                    "جوال",
                    phone,
                )
                for phone in extract_sa_phones(page_text)[:2]
            ]
        )

    if not candidate.has_social:
        for field, links in extract_social_links(
            page_text, company_domain=candidate.domain
        ).items():
            rows.extend(
                [
                    EnrichmentRow(
                        candidate.account_id,
                        candidate.company_name,
                        candidate.domain,
                        candidate.tier_bucket,
                        field,
                        link,
                    )
                    for link in links[:2]
                ]
            )

    return rows


def issue_from_result(
    candidate: EnrichmentCandidate,
    result: AgentReachResult | None,
    *,
    reason: str | None = None,
) -> QualityIssue:
    if reason:
        message = reason
    elif result is None:
        message = "لم تتم قراءة الموقع عبر Agent Reach."
    elif result.error:
        message = f"تعذر قراءة الموقع عبر Agent Reach: {result.error[:300]}"
    else:
        message = (
            "قرأ Agent Reach الموقع لكن لم يعثر على جوال سعودي "
            "أو روابط سوشيال قابلة للاستخدام."
        )
    return QualityIssue(candidate.account_id, candidate.company_name, candidate.domain, message)


async def enrich_candidate(
    candidate: EnrichmentCandidate,
    service: AgentReachService,
) -> tuple[list[EnrichmentRow], QualityIssue | None]:
    """Read one company homepage through Agent Reach and extract candidates."""
    result = await service.read_web_page(f"https://{candidate.domain}")
    if not result.success or not result.data:
        return [], issue_from_result(candidate, result)

    page_text = str(result.data)
    if len(page_text.strip()) < MIN_PAGE_TEXT_CHARS:
        return [], issue_from_result(
            candidate,
            result,
            reason="محتوى الموقع قصير جداً أو غير قابل للتحليل.",
        )

    try:
        rows = build_enrichment_rows(candidate, page_text)
    except ValueError as exc:
        return [], issue_from_result(
            candidate,
            result,
            reason=f"تعذر تحليل محتوى الموقع: {type(exc).__name__}: {str(exc)[:200]}",
        )
    if not rows:
        return [], issue_from_result(candidate, result)
    return rows, None
