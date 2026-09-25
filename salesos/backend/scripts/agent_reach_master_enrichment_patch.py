"""Build a review-only Agent Reach patch against a copy of the master CSV.

Read-only inputs: the cleaned Agent Reach workbook + the original master CSV.
Writes only under salesos/backend/outputs. Never overwrites original master
fields, the Downloads master CSV, the workbook, or any database.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse, urlunparse

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.modules.agent_reach.contact_enrichment import (  # noqa: E402
    ISSUES_SHEET,
    LINKEDIN_ALLOWED_FIRST_SEGMENTS,
    MARKDOWN_IMAGE_ARTIFACT,
    READY_SHEET,
    TWITTER_RESERVED_FIRST_SEGMENTS,
    _clean_url,
    _has_concatenated_http,
    _normalize_linkedin_url,
    _normalize_phone,
    _normalize_twitter_profile_url,
    _normalize_whatsapp_url,
    extract_sa_phones,
    sanitize_ready_url_value,
)

VALUE_DELIMITER = "|"
SOURCE_METHOD = "agent_reach_review_workbook"
BOTH_LIST_PREVIEW = 200

# Mapped from prior READY_FOR_REVIEW / QUALITY_REVIEW_REQUIRED labels.
STATUS_READY = "PENDING_REVIEW"
STATUS_READY_WITH_ISSUES = "READY_WITH_ISSUES"
STATUS_QUALITY = "QUALITY_BACKLOG"
STATUS_PENDING_REVIEW = STATUS_READY
STATUS_QUALITY_BACKLOG = STATUS_QUALITY

PATCH_HEADERS = [
    "Master Account ID",
    "company_name",
    "source_domain",
    "tier",
    "field_type",
    "value",
    "source_method",
    "review_status",
    "quality_flag",
    "notes",
]
QUALITY_HEADERS = [
    "Master Account ID",
    "company_name",
    "domain",
    "issue",
    "issue_category",
    "recommended_action",
]
ADDITIVE_COLUMNS = [
    "AgentReach_Phones",
    "AgentReach_WhatsApp",
    "AgentReach_LinkedIn",
    "AgentReach_Instagram",
    "AgentReach_X",
    "AgentReach_Facebook",
    "AgentReach_TikTok",
    "AgentReach_YouTube",
    "AgentReach_Snapchat",
    "AgentReach_Enrichment_Status",
    "AgentReach_Ready_Row_Count",
    "AgentReach_Quality_Issue",
    "Domain_Quality_Status",
    "Domain_Quality_Recommended_Action",
]
PROTECTED_MASTER_FIELDS = frozenset({
    "Primary_Phone",
    "Primary_Email",
    "Primary_Domain",
    "CR_Numbers",
    "Canonical_Company_Name",
})

PHONE_FIELDS = frozenset({"جوال", "هاتف"})
WHATSAPP_FIELDS = frozenset({"واتساب"})
LINKEDIN_FIELDS = frozenset({"لينكدإن", "لينكدان"})
TWITTER_FIELDS = frozenset({"تويتر", "تويتر/X", "تويتر/اكس"})

_FIELD_ALIASES = {
    "جوال": "جوال",
    "هاتف": "هاتف",
    "واتساب": "واتساب",
    "لينكدإن": "لينكدإن",
    "لينكدان": "لينكدإن",
    "إنستغرام": "انستقرام",
    "انستغرام": "انستقرام",
    "انستقرام": "انستقرام",
    "تويتر": "تويتر/X",
    "تويتر/x": "تويتر/X",
    "تويتر/اكس": "تويتر/X",
    "فيسبوك": "فيسبوك",
    "تيك توك": "تيك توك",
    "يوتيوب": "يوتيوب",
    "سناب": "سناب شات",
    "سناب شات": "سناب شات",
}
_FIELD_TO_COLUMN = {
    "جوال": "AgentReach_Phones",
    "هاتف": "AgentReach_Phones",
    "واتساب": "AgentReach_WhatsApp",
    "لينكدإن": "AgentReach_LinkedIn",
    "انستقرام": "AgentReach_Instagram",
    "تويتر/X": "AgentReach_X",
    "فيسبوك": "AgentReach_Facebook",
    "تيك توك": "AgentReach_TikTok",
    "يوتيوب": "AgentReach_YouTube",
    "سناب شات": "AgentReach_Snapchat",
}
LINKEDIN_STRIP_TAILS = frozenset({
    "about",
    "admin",
    "details",
    "mycompany",
    "people",
    "posts",
})
SLUG_RE = re.compile(r"^[a-zA-Z0-9][a-zA-Z0-9._-]{1,80}$")
LINKEDIN_COMPANY_SLUG_PARTS = 2
SA_DIGITS_WITH_TRUNK_ZERO = 13
SA_INTL_MOBILE_DIGITS = 12
SA_DOUBLE_ZERO_TRUNK_MIN = 14

# First matching needle wins. Needles are matched case-insensitively.
_ISSUE_RULES: tuple[tuple[tuple[str, ...], str, str, str], ...] = (
    (
        ("خطأ إملائي", ".vom", "لاحقة غير حقيقية", "لاحقة غير صحيحة", "yopmail"),
        "typo_domain",
        "TYPO_DOMAIN",
        "Review Primary_Domain for a typo or unofficial lookalike; "
        "do not treat it as the company site.",
    ),
    (
        ("حكومي", "gov.sa", ".gov."),
        "government_domain",
        "GOVERNMENT_DOMAIN",
        "Remove the government domain from Primary_Domain; it is not a company website.",
    ),
    (
        ("مخترق", "seo hijacking", "كازينو", "رهانات"),
        "hijacked_or_compromised",
        "HIJACKED_OR_COMPROMISED",
        "Treat the domain as compromised or hijacked; do not apply extracted contacts.",
    ),
    (
        ("plesk", "استضافة", "parked", "معروض للبيع", "coming soon", "قيد الإنشاء"),
        "parked_or_hosting",
        "PARKED_OR_HOSTING",
        "Review expired, parked, or placeholder hosting; do not apply theme/host contacts.",
    ),
    (
        ("وهمي", "تجريبي", "555"),
        "placeholder_or_fake",
        "PLACEHOLDER_OR_FAKE",
        "Discard placeholder or fake values; do not copy them onto the master.",
    ),
    (
        (
            "غير سعودي",
            "رقم مصري",
            "+20",
            "+973",
            "+971",
            "كود دولة",
            "وليست سعودية",
            "وليست سعودي",
        ),
        "non_saudi_contact",
        "NON_SAUDI_CONTACT",
        "Do not apply foreign phone/social values to this Saudi master account.",
    ),
    (
        (
            "عدم تطابق",
            "لا يتطابق",
            "مختلفة تماماً",
            "مختلفة تماما",
            "شركة أخرى",
            "غير مرتبط",
            "غير مرتبطة",
        ),
        "company_domain_mismatch",
        "COMPANY_DOMAIN_MISMATCH",
        "Do not apply contacts from this domain; review company/domain identity first.",
    ),
    (
        ("عالمية", "عالمي", "الشركة الأم"),
        "global_parent_domain",
        "GLOBAL_PARENT_DOMAIN",
        "Confirm whether this is a global parent site versus the local Saudi account.",
    ),
    (
        ("timeout", "انتهاء المهلة", "انتهاء مهلة", "مهلة الاتصال"),
        "timeout",
        "TIMEOUT",
        "Retry later and confirm the domain is still live before any master update.",
    ),
    (
        ("فشل dns", "لا يعمل", "لا يستجيب", "تعذر الاتصال", "نطاق ميت"),
        "unreachable",
        "UNREACHABLE",
        "Review Primary_Domain DNS/liveness; do not invent replacement contacts.",
    ),
    (
        (
            "بدون بيانات اتصال",
            "لم يعثر",
            "لم يتم العثور على أي بيانات اتصال",
            "لا يحتوي",
        ),
        "no_contact_found",
        "NO_CONTACT_FOUND",
        "Keep the domain if valid; no public Saudi phone/social was found for review.",
    ),
    (
        ("فاسد", "غير صالح", "بريد إلكتروني مجاني", "بريد الكتروني مجاني"),
        "invalid_or_generic_domain",
        "INVALID_OR_GENERIC_DOMAIN",
        "Review Primary_Domain as invalid or generic; do not use it for enrichment.",
    ),
)


@dataclass(frozen=True)
class ReadyRecord:
    account_id: str
    company_name: str
    domain: str
    tier: str
    field: str
    value: str


@dataclass(frozen=True)
class IssueRecord:
    account_id: str
    company_name: str
    domain: str
    issue: str


@dataclass
class AccountPatch:
    account_id: str
    company_name: str = ""
    domain: str = ""
    tier: str = ""
    values_by_column: dict[str, list[str]] = field(default_factory=dict)
    ready_rows: list[ReadyRecord] = field(default_factory=list)
    issues: list[IssueRecord] = field(default_factory=list)
    duplicates_removed: int = 0
    unmatched: bool = False

    @property
    def has_ready(self) -> bool:
        return any(self.values_by_column.values())

    @property
    def has_issues(self) -> bool:
        return bool(self.issues)

    @property
    def status(self) -> str:
        return enrichment_status(self.has_ready, self.has_issues)

    @property
    def ready_row_count(self) -> int:
        return sum(len(values) for values in self.values_by_column.values())


def _timestamp() -> str:
    return datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")


def file_snapshot(path: Path) -> dict[str, str | int | float | None]:
    if not path.exists():
        return {
            "path": str(path),
            "exists": False,
            "size": None,
            "mtime_utc": None,
        }
    stat = path.stat()
    return {
        "path": str(path.resolve()),
        "exists": True,
        "size": stat.st_size,
        "mtime_utc": datetime.fromtimestamp(stat.st_mtime, UTC).isoformat(),
    }


def discover_master_id_column(fieldnames: Iterable[str] | None) -> str:
    names = [name for name in (fieldnames or []) if name]
    for name in names:
        if name.strip() == "Master Account ID":
            return name
    for name in names:
        lowered = name.strip().lower().replace("-", "_").replace(" ", "_")
        if lowered in {"master_account_id", "account_id"} or lowered.startswith("ma_"):
            return name
        if "master" in lowered and "id" in lowered:
            return name
    raise ValueError("Could not discover a Master Account ID column in the master CSV header.")


def canonical_field(raw: str) -> str:
    key = (raw or "").strip()
    if not key:
        return ""
    return _FIELD_ALIASES.get(key, _FIELD_ALIASES.get(key.lower(), key))


def field_to_column(raw_field: str) -> str | None:
    return _FIELD_TO_COLUMN.get(canonical_field(raw_field))


def enrichment_status(has_ready: bool, has_issues: bool) -> str:
    if has_ready and has_issues:
        return STATUS_READY_WITH_ISSUES
    if has_ready:
        return STATUS_READY
    if has_issues:
        return STATUS_QUALITY
    return ""


def quality_flag(status: str, value: str = "") -> str:
    if status in {STATUS_READY_WITH_ISSUES, STATUS_QUALITY}:
        return "needs_review"
    if MARKDOWN_IMAGE_ARTIFACT in (value or "") or _has_concatenated_http(value or ""):
        return "needs_review"
    if status == STATUS_READY:
        return "ok"
    return "needs_review"


def classify_quality_issue(issue: str) -> tuple[str, str, str]:
    text = (issue or "").strip()
    lowered = text.lower()
    for needles, category, status, action in _ISSUE_RULES:
        if any(needle.lower() in lowered for needle in needles):
            return category, status, action
    return (
        "other",
        "OTHER",
        "Human-review the quality-issue text before any master update.",
    )


def _dedupe_keep_order(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    unique: list[str] = []
    for item in values:
        key = item.lower().rstrip("/")
        if not item or key in seen:
            continue
        seen.add(key)
        unique.append(item)
    return unique


def recover_phone_values(raw: str) -> list[str]:
    """Return normalized Saudi phones; drop foreign/unparseable cells."""
    text = "" if raw is None else str(raw).strip()
    if not text:
        return []
    found = extract_sa_phones(text)
    if found:
        return _dedupe_keep_order(found)

    digits = re.sub(r"\D", "", text)
    candidates: list[str] = []
    if digits.startswith("009660") and len(digits) >= SA_DOUBLE_ZERO_TRUNK_MIN:
        candidates.append("966" + digits[6:])
    if digits.startswith("9660") and len(digits) == SA_DIGITS_WITH_TRUNK_ZERO:
        candidates.append("966" + digits[4:])
    if digits.startswith("966800") and len(digits) in {
        SA_INTL_MOBILE_DIGITS,
        SA_DIGITS_WITH_TRUNK_ZERO,
    }:
        candidates.append(digits[3:])
    if digits.startswith("996") and len(digits) == SA_INTL_MOBILE_DIGITS:
        candidates.append("966" + digits[3:])
    recovered: list[str] = []
    for candidate in candidates:
        normalized = _normalize_phone(candidate)
        if normalized:
            recovered.append(normalized)
    if recovered:
        return _dedupe_keep_order(recovered)

    parts = re.split(r"[/,|;]+", text)
    for part in parts:
        normalized = _normalize_phone(part.strip())
        if normalized:
            recovered.append(normalized)
        else:
            recovered.extend(extract_sa_phones(part))
    return _dedupe_keep_order(recovered)


def _whatsapp_url_candidates(raw: str) -> list[str]:
    text = raw.strip()
    candidates = [text, sanitize_ready_url_value(text), _clean_url(text)]
    lowered = text.lower()
    if lowered.startswith("wa.me"):
        candidates.append(f"https://{text.lstrip('/')}")
    match = re.search(r"(?:wa\.me/|phone=)(\+?0*9660?\d+)", text, flags=re.I)
    if match:
        candidates.append(f"https://wa.me/{match.group(1)}")
    return _dedupe_keep_order(candidates)


def recover_whatsapp_value(raw: str) -> str:
    text = "" if raw is None else str(raw).strip()
    if not text:
        return ""
    lowered = text.lower()
    if any(token in lowered for token in ("wa.link/", "iwtsp.com", "/channel/", "/message/")):
        return ""
    for candidate in _whatsapp_url_candidates(text):
        if "://" in candidate or candidate.lower().startswith("wa.me"):
            normalized = _normalize_whatsapp_url(candidate) or _normalize_whatsapp_url(
                _clean_url(candidate)
            )
            if normalized:
                return normalized
    phones = recover_phone_values(text)
    if not phones:
        return ""
    digits = phones[0].lstrip("+")
    if digits.startswith(("9200", "800")):
        digits = f"966{digits}"
    if not digits.startswith("966"):
        return ""
    rebuilt = f"https://wa.me/{digits}"
    return _normalize_whatsapp_url(rebuilt) or ""


def recover_linkedin_value(raw: str) -> str:
    text = "" if raw is None else str(raw).strip()
    if not text:
        return ""
    existing = _normalize_linkedin_url(text)
    if existing:
        return existing
    cleaned = sanitize_ready_url_value(text)
    cleaned = re.sub(r"[)\]].*$", "", cleaned).strip()
    if "linkedin.com" in cleaned.lower() and "://" not in cleaned:
        cleaned = f"https://{cleaned.lstrip('/')}"
    if "://" not in cleaned:
        lowered = cleaned.lower()
        if lowered.startswith(("company/", "in/", "school/", "showcase/")):
            cleaned = f"https://www.linkedin.com/{cleaned.lstrip('/')}"
        elif SLUG_RE.fullmatch(cleaned):
            cleaned = f"https://www.linkedin.com/company/{cleaned}"
        else:
            return ""
    parsed = urlparse(cleaned)
    if parsed is None:
        return ""
    parts = [part for part in (parsed.path or "").strip("/").split("/") if part]
    while parts and parts[-1].lower().split("?")[0] in LINKEDIN_STRIP_TAILS:
        parts.pop()
    if (
        parts
        and parts[0].lower() in LINKEDIN_ALLOWED_FIRST_SEGMENTS
        and len(parts) > LINKEDIN_COMPANY_SLUG_PARTS
    ):
        parts = parts[:2]
    host = parsed.netloc or "www.linkedin.com"
    rebuilt = urlunparse((parsed.scheme or "https", host, "/" + "/".join(parts), "", "", ""))
    return _normalize_linkedin_url(rebuilt) or ""


def recover_twitter_value(raw: str) -> str:
    text = "" if raw is None else str(raw).strip()
    if not text:
        return ""
    cleaned = sanitize_ready_url_value(text)
    normalized = _normalize_twitter_profile_url(cleaned)
    if normalized:
        return normalized
    parsed = urlparse(cleaned if "://" in cleaned else f"https://{cleaned.lstrip('/')}")
    parts = [part for part in (parsed.path or "").strip("/").split("/") if part]
    if not parts:
        return ""
    handle = parts[0]
    if handle.lower() in TWITTER_RESERVED_FIRST_SEGMENTS:
        return ""
    host = parsed.netloc or "x.com"
    rebuilt = urlunparse((parsed.scheme or "https", host, f"/{handle}", "", "", ""))
    return _normalize_twitter_profile_url(rebuilt) or ""


def recover_ready_values(raw_field: str, value: str) -> list[str]:  # noqa: PLR0911
    """Return guard-passing ready values, or empty if the cell is an FP."""
    raw = "" if value is None else str(value).strip()
    if not raw:
        return []
    field = canonical_field(raw_field)
    if field in PHONE_FIELDS:
        return recover_phone_values(raw)
    if field in WHATSAPP_FIELDS:
        recovered = recover_whatsapp_value(raw)
        return [recovered] if recovered else []
    if field in LINKEDIN_FIELDS:
        recovered = recover_linkedin_value(raw)
        return [recovered] if recovered else []
    if field in TWITTER_FIELDS:
        recovered = recover_twitter_value(raw)
        return [recovered] if recovered else []
    cleaned = sanitize_ready_url_value(raw)
    if _has_concatenated_http(cleaned) or MARKDOWN_IMAGE_ARTIFACT in cleaned:
        cleaned = _clean_url(cleaned)
    if not cleaned:
        return []
    if cleaned.lower().startswith(("http://", "https://")):
        parsed = urlparse(cleaned)
        query = parse_qs(parsed.query) if parsed else {}
        if "text" in query:
            return []
    return [cleaned]


def normalize_patch_value(raw_field: str, value: str) -> str:
    recovered = recover_ready_values(raw_field, value)
    return recovered[0] if recovered else ""


def value_dedupe_key(raw_field: str, value: str) -> str:
    field = canonical_field(raw_field)
    if field in PHONE_FIELDS:
        return (_normalize_phone(value) or value).lower()
    return value.lower().rstrip("/")


def join_values(values: Iterable[str]) -> str:
    return VALUE_DELIMITER.join(item for item in values if item)


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _row_cell(row: tuple[object, ...] | None, index: int) -> str:
    if not row or len(row) <= index:
        return ""
    return _cell(row[index])


def load_workbook_records(workbook_path: Path) -> tuple[list[ReadyRecord], list[IssueRecord]]:
    from openpyxl import load_workbook

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ready: list[ReadyRecord] = []
    issues: list[IssueRecord] = []
    if READY_SHEET in wb.sheetnames:
        for row in wb[READY_SHEET].iter_rows(min_row=2, values_only=True):
            account_id = _row_cell(row, 0)
            if not account_id:
                continue
            ready.append(
                ReadyRecord(
                    account_id=account_id,
                    company_name=_row_cell(row, 1),
                    domain=_row_cell(row, 2),
                    tier=_row_cell(row, 3),
                    field=_row_cell(row, 4),
                    value=_row_cell(row, 5),
                )
            )
    if ISSUES_SHEET in wb.sheetnames:
        for row in wb[ISSUES_SHEET].iter_rows(min_row=2, values_only=True):
            account_id = _row_cell(row, 0)
            if not account_id:
                continue
            issues.append(
                IssueRecord(
                    account_id=account_id,
                    company_name=_row_cell(row, 1),
                    domain=_row_cell(row, 2),
                    issue=_row_cell(row, 3),
                )
            )
    wb.close()
    return ready, issues


def build_account_patches(
    ready_rows: Iterable[ReadyRecord],
    issues: Iterable[IssueRecord],
) -> dict[str, AccountPatch]:
    patches: dict[str, AccountPatch] = {}

    def _patch(account_id: str) -> AccountPatch:
        return patches.setdefault(account_id, AccountPatch(account_id=account_id))

    for row in ready_rows:
        patch = _patch(row.account_id)
        if not patch.company_name:
            patch.company_name = row.company_name
        if not patch.domain:
            patch.domain = row.domain
        if not patch.tier:
            patch.tier = row.tier
        patch.ready_rows.append(row)
        column = field_to_column(row.field)
        if column is None:
            continue
        existing = patch.values_by_column.setdefault(column, [])
        recovered = recover_ready_values(row.field, row.value)
        if not recovered:
            continue
        for normalized in recovered:
            key = value_dedupe_key(row.field, normalized)
            if any(value_dedupe_key(row.field, item) == key for item in existing):
                patch.duplicates_removed += 1
                continue
            existing.append(normalized)

    for issue in issues:
        patch = _patch(issue.account_id)
        if not patch.company_name:
            patch.company_name = issue.company_name
        if not patch.domain:
            patch.domain = issue.domain
        patch.issues.append(issue)

    return patches


def _issue_summary(patch: AccountPatch) -> tuple[str, str, str]:
    if not patch.issues:
        return "", "", ""
    categories: list[str] = []
    statuses: list[str] = []
    actions: list[str] = []
    texts: list[str] = []
    for issue in patch.issues:
        category, status, action = classify_quality_issue(issue.issue)
        if category not in categories:
            categories.append(category)
        if status not in statuses:
            statuses.append(status)
        if action not in actions:
            actions.append(action)
        if issue.issue and issue.issue not in texts:
            texts.append(issue.issue)
    return (
        join_values(texts),
        join_values(statuses),
        join_values(actions),
    )


def additive_columns_for(patch: AccountPatch | None) -> dict[str, str]:
    out = {column: "" for column in ADDITIVE_COLUMNS}
    if patch is None:
        return out
    for column, values in patch.values_by_column.items():
        out[column] = join_values(values)
    out["AgentReach_Enrichment_Status"] = patch.status
    out["AgentReach_Ready_Row_Count"] = str(patch.ready_row_count) if patch.has_ready else ""
    issue_text, domain_status, domain_action = _issue_summary(patch)
    out["AgentReach_Quality_Issue"] = issue_text
    out["Domain_Quality_Status"] = domain_status
    out["Domain_Quality_Recommended_Action"] = domain_action
    return out


def build_patch_rows(
    patches: dict[str, AccountPatch],
    *,
    master_names: dict[str, str] | None = None,
) -> list[list[str]]:
    rows: list[list[str]] = []
    for account_id in sorted(patches):
        patch = patches[account_id]
        company = (master_names or {}).get(account_id) or patch.company_name
        status = patch.status
        if not patch.ready_rows:
            continue
        emitted = False
        seen_keys: set[tuple[str, str]] = set()
        for ready in patch.ready_rows:
            column = field_to_column(ready.field)
            recovered = recover_ready_values(ready.field, ready.value)
            if not recovered:
                continue
            for normalized in recovered:
                field_key = column or canonical_field(ready.field)
                key = (field_key, value_dedupe_key(ready.field, normalized))
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                notes = []
                if patch.unmatched:
                    notes.append("unmatched_master_id")
                if column is None:
                    notes.append("unmapped_field")
                rows.append(
                    [
                        account_id,
                        company,
                        ready.domain or patch.domain,
                        ready.tier or patch.tier,
                        canonical_field(ready.field) or ready.field,
                        normalized,
                        SOURCE_METHOD,
                        status,
                        quality_flag(status, normalized),
                        ";".join(notes),
                    ]
                )
                emitted = True
        if not emitted and patch.unmatched:
            rows.append(
                [
                    account_id,
                    company,
                    patch.domain,
                    patch.tier,
                    "",
                    "",
                    SOURCE_METHOD,
                    status,
                    quality_flag(status),
                    "unmatched_master_id",
                ]
            )
    return rows


def build_quality_rows(
    patches: dict[str, AccountPatch],
    *,
    master_names: dict[str, str] | None = None,
) -> list[list[str]]:
    rows: list[list[str]] = []
    for account_id in sorted(patches):
        patch = patches[account_id]
        company = (master_names or {}).get(account_id) or patch.company_name
        for issue in patch.issues:
            category, _status, action = classify_quality_issue(issue.issue)
            rows.append(
                [
                    account_id,
                    company,
                    issue.domain or patch.domain,
                    issue.issue,
                    category,
                    action,
                ]
            )
    return rows


def _write_csv(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def write_enriched_master(
    master_csv_path: Path,
    output_path: Path,
    patches: dict[str, AccountPatch],
    *,
    id_column: str | None = None,
) -> dict[str, object]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    matched: set[str] = set()
    master_ids: set[str] = set()
    master_names: dict[str, str] = {}
    total_rows = 0
    original_fields: list[str] = []

    with master_csv_path.open("r", encoding="utf-8-sig", newline="") as src:
        reader = csv.DictReader(src)
        original_fields = list(reader.fieldnames or [])
        resolved_id = id_column or discover_master_id_column(original_fields)
        collision = [col for col in ADDITIVE_COLUMNS if col in original_fields]
        if collision:
            raise ValueError(
                "Refusing to write additive columns that already exist on master: "
                + ", ".join(collision)
            )
        fieldnames = original_fields + ADDITIVE_COLUMNS
        with output_path.open("w", encoding="utf-8-sig", newline="") as dst:
            writer = csv.DictWriter(dst, fieldnames=fieldnames, extrasaction="raise")
            writer.writeheader()
            for row in reader:
                total_rows += 1
                account_id = (row.get(resolved_id) or "").strip()
                if account_id:
                    master_ids.add(account_id)
                    name = (row.get("Canonical_Company_Name") or "").strip()
                    if name:
                        master_names[account_id] = name
                out = {key: row.get(key, "") for key in original_fields}
                patch = patches.get(account_id) if account_id else None
                if patch is not None:
                    matched.add(account_id)
                    patch.unmatched = False
                    if name := master_names.get(account_id):
                        patch.company_name = name
                out.update(additive_columns_for(patch))
                writer.writerow(out)

    for account_id, patch in patches.items():
        if account_id not in master_ids:
            patch.unmatched = True

    return {
        "id_column": resolved_id,
        "total_master_rows": total_rows,
        "original_field_count": len(original_fields),
        "original_fields": original_fields,
        "matched_ids": matched,
        "master_ids": master_ids,
        "master_names": master_names,
    }


def _md_list(items: Iterable[str], *, empty: str = "none") -> str:
    values = list(items)
    if not values:
        return empty
    return "\n".join(f"- `{item}`" for item in values)


def render_report(stats: dict[str, object]) -> str:
    unmatched = list(stats["unmatched_ids"])
    both = list(stats["accounts_with_ready_and_issues"])
    lines = [
        "# Master Enrichment Patch Report",
        "",
        "Review-only patch showing how Agent Reach workbook values would be",
        "reflected onto a **copy** of the master CSV. Additive columns only.",
        "",
        f"Generated at (UTC): `{stats['generated_at']}`",
        "",
        "## Inputs (read-only)",
        "",
        f"- Workbook: `{stats['workbook']}`",
        f"- Master CSV: `{stats['master_csv']}`",
        "",
        "## Outputs (writes only under salesos/backend/outputs)",
        "",
        f"- Directory: `{stats['output_dir']}`",
        f"- `{stats['patch_csv']}`",
        f"- `{stats['quality_csv']}`",
        f"- `{stats['enriched_csv']}`",
        f"- Report: `{stats['report_path']}`",
        "",
        "## Counts",
        "",
        f"- Total master rows: **{stats['total_master_rows']}**",
        f"- Master ID column: `{stats['id_column']}`",
        f"- Agent Reach account IDs (ready + issues): **{stats['agent_reach_account_ids']}**",
        f"- Matched Agent Reach account IDs: **{stats['matched_account_ids']}**",
        f"- Unmatched Agent Reach account IDs: **{len(unmatched)}**",
        f"- Ready rows in workbook: **{stats['ready_rows_in_workbook']}**",
        f"- Ready rows applied into additive columns (after dedupe): "
        f"**{stats['ready_rows_applied']}**",
        f"- Unique accounts enriched (ready values on a matched master row): "
        f"**{stats['unique_accounts_enriched']}**",
        f"- Quality issue rows: **{stats['quality_issue_count']}**",
        f"- Accounts with both ready rows and quality issues: **{len(both)}**",
        f"- Duplicate values removed: **{stats['duplicates_removed']}**",
        "",
        "## Counts by tier (ready workbook rows, after field mapping)",
        "",
    ]
    tier_counts: Counter[str] = stats["tier_counts"]  # type: ignore[assignment]
    if tier_counts:
        for tier, count in sorted(tier_counts.items()):
            lines.append(f"- `{tier or 'UNKNOWN'}`: {count}")
    else:
        lines.append("- none")
    lines.extend(["", "## Counts by field (ready rows applied after dedupe)", ""])
    field_counts: Counter[str] = stats["field_counts"]  # type: ignore[assignment]
    if field_counts:
        for field_name, count in sorted(field_counts.items()):
            lines.append(f"- `{field_name}`: {count}")
    else:
        lines.append("- none")
    lines.extend(
        [
            "",
            "## Duplicate value checks",
            "",
            "- Dedupe key: lowercase URL without trailing `/`, or normalized Saudi phone.",
            f"- Multi-value delimiter in additive columns: `{VALUE_DELIMITER}` (no spaces).",
            "- Master `All_Phones` / `All_Domains` already use `; ` — Agent Reach additive",
            "  columns use `|` so URL/phone lists stay unambiguous.",
            f"- Duplicate values collapsed: **{stats['duplicates_removed']}**",
            "",
            "## Accounts with both ready rows and quality issues",
            "",
            _md_list(both[:BOTH_LIST_PREVIEW], empty="none"),
        ]
    )
    if len(both) > BOTH_LIST_PREVIEW:
        leftover = len(both) - BOTH_LIST_PREVIEW
        lines.append(f"- … {leftover} more (see patch CSVs)")
    lines.extend(
        [
            "",
            "## Unmatched Agent Reach account IDs",
            "",
            "These IDs appear in the workbook but not in the master CSV.",
            "They are listed here and kept in the review patch CSVs only.",
            "No master rows were invented.",
            "",
            _md_list(unmatched, empty="none"),
            "",
            "## Safety confirmation",
            "",
            "Original master CSV (Downloads) — before:",
            "",
            f"- path: `{stats['master_before']['path']}`",
            f"- exists: `{stats['master_before']['exists']}`",
            f"- size: `{stats['master_before']['size']}`",
            f"- mtime_utc: `{stats['master_before']['mtime_utc']}`",
            "",
            "Original master CSV (Downloads) — after:",
            "",
            f"- path: `{stats['master_after']['path']}`",
            f"- exists: `{stats['master_after']['exists']}`",
            f"- size: `{stats['master_after']['size']}`",
            f"- mtime_utc: `{stats['master_after']['mtime_utc']}`",
            f"- unchanged: **{stats['master_unchanged']}**",
            "",
            "Closure workbook (read-only) — before / after:",
            "",
            f"- before size/mtime: `{stats['workbook_before']['size']}` / "
            f"`{stats['workbook_before']['mtime_utc']}`",
            f"- after size/mtime: `{stats['workbook_after']['size']}` / "
            f"`{stats['workbook_after']['mtime_utc']}`",
            f"- unchanged: **{stats['workbook_unchanged']}**",
            "",
            "Downloads Agent Reach workbook:",
            "",
            f"- `{stats['downloads_workbook_note']}`",
            "",
            "- Production DB: not opened, not written.",
            "- Master data DB tables: not opened, not written.",
            f"- All write paths under `salesos/backend/outputs`: **{stats['outputs_only']}**",
            "",
            "REVIEW-ONLY MASTER PATCH CREATED",
            "NO ORIGINAL MASTER OVERWRITE",
            "PRODUCTION NOT APPROVED",
            "",
        ]
    )
    return "\n".join(lines)


def run_patch(
    *,
    workbook_path: Path,
    master_csv_path: Path,
    output_dir: Path,
    report_path: Path,
    downloads_workbook: Path | None = None,
) -> dict[str, object]:
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)

    master_before = file_snapshot(master_csv_path)
    workbook_before = file_snapshot(workbook_path)
    downloads_note = (
        "no Agent Reach workbook found/used under Downloads "
        "(input workbook is the outputs closure_cleaned copy)."
    )
    if downloads_workbook is not None:
        snap = file_snapshot(downloads_workbook)
        downloads_note = (
            f"observed `{snap['path']}` size={snap['size']} mtime_utc={snap['mtime_utc']} "
            "(not used as a write target)."
        )

    ready_rows, issues = load_workbook_records(workbook_path)
    patches = build_account_patches(ready_rows, issues)

    enriched_csv = output_dir / "01_Master_Accounts_enriched_review_only.csv"
    apply_stats = write_enriched_master(master_csv_path, enriched_csv, patches)
    master_names = apply_stats["master_names"]  # type: ignore[assignment]

    patch_rows = build_patch_rows(patches, master_names=master_names)
    quality_rows = build_quality_rows(patches, master_names=master_names)
    patch_csv = output_dir / "master_enrichment_patch.csv"
    quality_csv = output_dir / "master_quality_issues_patch.csv"
    _write_csv(patch_csv, PATCH_HEADERS, patch_rows)
    _write_csv(quality_csv, QUALITY_HEADERS, quality_rows)

    unmatched_ids = sorted(
        account_id for account_id, patch in patches.items() if patch.unmatched
    )
    matched_ids = sorted(apply_stats["matched_ids"])  # type: ignore[arg-type]
    both = sorted(
        account_id
        for account_id, patch in patches.items()
        if patch.has_ready and patch.has_issues and not patch.unmatched
    )
    unique_enriched = sum(
        1 for account_id, patch in patches.items() if patch.has_ready and not patch.unmatched
    )
    ready_applied = sum(
        patch.ready_row_count for patch in patches.values() if not patch.unmatched
    )
    duplicates_removed = sum(patch.duplicates_removed for patch in patches.values())
    tier_counts: Counter[str] = Counter()
    field_counts: Counter[str] = Counter()
    for patch in patches.values():
        if patch.unmatched:
            continue
        for ready in patch.ready_rows:
            if field_to_column(ready.field):
                tier_counts[ready.tier or patch.tier] += 1
        for column, values in patch.values_by_column.items():
            field_counts[column] += len(values)

    master_after = file_snapshot(master_csv_path)
    workbook_after = file_snapshot(workbook_path)
    report_copy = output_dir / "MASTER_ENRICHMENT_PATCH_REPORT.md"
    write_paths = [patch_csv, quality_csv, enriched_csv, report_path, report_copy]
    outputs_root = (BACKEND_ROOT / "outputs").resolve()
    outputs_only = all(
        path.resolve() == outputs_root or outputs_root in path.resolve().parents
        for path in write_paths
    )

    stats: dict[str, object] = {
        "generated_at": datetime.now(UTC).isoformat(),
        "workbook": str(workbook_path.resolve()),
        "master_csv": str(master_csv_path.resolve()),
        "output_dir": str(output_dir.resolve()),
        "patch_csv": str(patch_csv.resolve()),
        "quality_csv": str(quality_csv.resolve()),
        "enriched_csv": str(enriched_csv.resolve()),
        "report_path": str(report_path.resolve()),
        "id_column": apply_stats["id_column"],
        "total_master_rows": apply_stats["total_master_rows"],
        "agent_reach_account_ids": len(patches),
        "matched_account_ids": len(matched_ids),
        "unmatched_ids": unmatched_ids,
        "ready_rows_in_workbook": len(ready_rows),
        "ready_rows_applied": ready_applied,
        "unique_accounts_enriched": unique_enriched,
        "quality_issue_count": len(quality_rows),
        "accounts_with_ready_and_issues": both,
        "duplicates_removed": duplicates_removed,
        "tier_counts": tier_counts,
        "field_counts": field_counts,
        "master_before": master_before,
        "master_after": master_after,
        "master_unchanged": master_before == master_after,
        "workbook_before": workbook_before,
        "workbook_after": workbook_after,
        "workbook_unchanged": workbook_before == workbook_after,
        "downloads_workbook_note": downloads_note,
        "outputs_only": outputs_only,
    }
    report = render_report(stats)
    report_path.write_text(report, encoding="utf-8")
    (output_dir / "MASTER_ENRICHMENT_PATCH_REPORT.md").write_text(report, encoding="utf-8")
    return stats


def _default_workbook() -> Path:
    return (
        BACKEND_ROOT
        / "outputs"
        / "agent_reach_contact_enrichment"
        / "20260906T113207Z_closure_cleaned"
        / "14_Website_Phone_Social_Enrichment_AgentReach.xlsx"
    )


def _default_master() -> Path:
    return Path.home() / "Downloads" / "MUHIDE_extracted" / "01_Master_Accounts.csv"


def _default_report() -> Path:
    return (
        BACKEND_ROOT
        / "outputs"
        / "agent_reach_contact_enrichment"
        / "MASTER_ENRICHMENT_PATCH_REPORT.md"
    )


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Create a review-only Agent Reach master enrichment patch."
    )
    parser.add_argument("--workbook", type=Path, default=_default_workbook())
    parser.add_argument("--master-csv", type=Path, default=_default_master())
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=(
            BACKEND_ROOT
            / "outputs"
            / "agent_reach_contact_enrichment"
            / f"{_timestamp()}_master_patch"
        ),
    )
    parser.add_argument("--report-path", type=Path, default=_default_report())
    parser.add_argument("--downloads-workbook", type=Path, default=None)
    return parser


def main() -> int:
    args = _build_parser().parse_args()
    stats = run_patch(
        workbook_path=args.workbook,
        master_csv_path=args.master_csv,
        output_dir=args.output_dir,
        report_path=args.report_path,
        downloads_workbook=args.downloads_workbook,
    )
    print(f"output_dir={stats['output_dir']}")
    print(f"total_master_rows={stats['total_master_rows']}")
    print(f"matched_account_ids={stats['matched_account_ids']}")
    print(f"unmatched_account_ids={len(stats['unmatched_ids'])}")  # type: ignore[arg-type]
    print(f"ready_rows_applied={stats['ready_rows_applied']}")
    print(f"unique_accounts_enriched={stats['unique_accounts_enriched']}")
    print(f"quality_issue_count={stats['quality_issue_count']}")
    print(f"master_unchanged={stats['master_unchanged']}")
    print(f"outputs_only={stats['outputs_only']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
