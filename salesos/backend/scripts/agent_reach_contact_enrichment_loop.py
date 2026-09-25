"""Continue MUHIDE website/phone/social enrichment through Agent Reach.

The script is intentionally review-only:
- reads MUHIDE master CSV + the in-progress workbook;
- skips accounts already present in the workbook;
- fetches one official company homepage through AgentReachService;
- writes new candidate rows and quality issues to a fresh output folder.

It never writes to SalesOS production data or to the original workbook.
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
import shutil
import sys
from copy import copy
from datetime import UTC, datetime
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.modules.agent_reach.contact_enrichment import (  # noqa: E402
    ISSUES_SHEET,
    MARKDOWN_IMAGE_ARTIFACT,
    QUEUE_HEADERS,
    READY_SHEET,
    EnrichmentCandidate,
    EnrichmentRow,
    QualityIssue,
    enrich_candidate,
    load_master_candidates,
    load_processed_accounts,
    load_queue_candidates,
    queue_row_from_candidate,
    sanitize_ready_url_value,
)
from app.modules.agent_reach.service import AgentReachService  # noqa: E402

READY_HEADERS = ["رقم الحساب", "اسم الشركة", "النطاق", "الفئة", "الحقل", "القيمة/الرابط"]
ISSUE_HEADERS = ["رقم الحساب", "اسم الشركة", "النطاق المسجل", "المشكلة"]
READY_VALUE_COL = READY_HEADERS.index("القيمة/الرابط") + 1


def _default_downloads_path(filename: str) -> Path:
    return Path.home() / "Downloads" / filename


def _timestamp() -> str:
    return datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")


def _parse_tiers(value: str) -> set[str]:
    return {item.strip().upper() for item in value.split(",") if item.strip()}


def _write_csv(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def _append_checkpoint(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(payload, ensure_ascii=False, default=str) + "\n")


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)


def _ensure_sheet(wb, name: str, headers: list[str]):
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(name)
        ws.append(headers)
    if ws.max_row < 1:
        ws.append(headers)
    return ws


def _update_first_table_ref(ws) -> None:
    tables = list(ws.tables.values())
    if not tables:
        return
    table = tables[0]
    end_col = table.ref.split(":")[1]
    end_col_letters = "".join(ch for ch in end_col if ch.isalpha())
    table.ref = f"A1:{end_col_letters}{ws.max_row}"


def _append_to_workbook(
    source_workbook: Path,
    output_workbook: Path,
    ready_rows: list[EnrichmentRow],
    issues: list[QualityIssue],
) -> None:
    from openpyxl import load_workbook

    shutil.copy2(source_workbook, output_workbook)
    wb = load_workbook(output_workbook)

    ready_ws = _ensure_sheet(wb, READY_SHEET, READY_HEADERS)
    issue_ws = _ensure_sheet(wb, ISSUES_SHEET, ISSUE_HEADERS)

    ready_style_row = ready_ws.max_row if ready_ws.max_row > 1 else 1
    for row in ready_rows:
        next_row = ready_ws.max_row + 1
        _copy_row_style(ready_ws, ready_style_row, next_row, len(READY_HEADERS))
        for col, value in enumerate(row.as_csv_row(), start=1):
            ready_ws.cell(next_row, col).value = value

    issue_style_row = issue_ws.max_row if issue_ws.max_row > 1 else 1
    for issue in issues:
        next_row = issue_ws.max_row + 1
        _copy_row_style(issue_ws, issue_style_row, next_row, len(ISSUE_HEADERS))
        for col, value in enumerate(issue.as_csv_row(), start=1):
            issue_ws.cell(next_row, col).value = value

    _update_first_table_ref(ready_ws)
    _update_first_table_ref(issue_ws)
    wb.save(output_workbook)


def _sheet_data_rows(ws) -> int:
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(str(cell or "").strip() for cell in row):
            count += 1
    return count


def sanitize_existing_workbook(source_workbook: Path, output_workbook: Path) -> dict:
    """Copy a review workbook and strip ``![Image`` leaks from ready URL cells.

    Quality-issue rows are left unchanged. The source file is never overwritten.
    """
    from openpyxl import load_workbook

    if source_workbook.resolve() == output_workbook.resolve():
        raise ValueError("sanitize must write a new workbook; refusing in-place overwrite")

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_workbook, output_workbook)
    wb = load_workbook(output_workbook)

    ready_rows = 0
    issue_rows = 0
    artifacts_before = 0
    artifacts_after = 0
    changed: list[dict[str, str | int]] = []

    if READY_SHEET in wb.sheetnames:
        ready_ws = wb[READY_SHEET]
        ready_rows = _sheet_data_rows(ready_ws)
        for row_idx in range(2, ready_ws.max_row + 1):
            cell = ready_ws.cell(row_idx, READY_VALUE_COL)
            raw = "" if cell.value is None else str(cell.value)
            if MARKDOWN_IMAGE_ARTIFACT in raw:
                artifacts_before += 1
            cleaned = sanitize_ready_url_value(raw)
            if cleaned != raw:
                cell.value = cleaned
                changed.append(
                    {
                        "excel_row": row_idx,
                        "account_id": str(ready_ws.cell(row_idx, 1).value or ""),
                        "field": str(ready_ws.cell(row_idx, 5).value or ""),
                        "before": raw,
                        "after": cleaned,
                    }
                )
            if MARKDOWN_IMAGE_ARTIFACT in str(cell.value or ""):
                artifacts_after += 1

    if ISSUES_SHEET in wb.sheetnames:
        issue_rows = _sheet_data_rows(wb[ISSUES_SHEET])

    wb.save(output_workbook)
    return {
        "source": str(source_workbook.resolve()),
        "output": str(output_workbook.resolve()),
        "ready_rows": ready_rows,
        "issue_rows": issue_rows,
        "ready_urls_cleaned": len(changed),
        "image_artifacts_ready_before": artifacts_before,
        "image_artifacts_ready_after": artifacts_after,
        "changed": changed,
        "review_only": True,
    }


def worker_writes_workbook(*, queue_csv: Path | None, no_workbook: bool) -> bool:
    """Workers never write a workbook; only a single-agent run may copy one."""
    return queue_csv is None and not no_workbook


def _persist_worker_outputs(
    output_dir: Path,
    ready_rows: list[EnrichmentRow],
    issues: list[QualityIssue],
    stats: dict,
    workbook_path: Path | None = None,
) -> dict:
    """Flush CSV + summary after each attempt so a killed worker keeps partial work."""
    _write_csv(
        output_dir / "agent_reach_ready_rows.csv",
        READY_HEADERS,
        [row.as_csv_row() for row in ready_rows],
    )
    _write_csv(
        output_dir / "agent_reach_quality_issues.csv",
        ISSUE_HEADERS,
        [issue.as_csv_row() for issue in issues],
    )
    summary = {
        "mode": stats.get("mode", "single"),
        "assigned": stats["assigned"],
        "attempted": stats["attempted"],
        "processed": stats["attempted"],
        "ready_rows": len(ready_rows),
        "issue_rows": len(issues),
        "incomplete": stats["incomplete"],
        "workbook": str(workbook_path.resolve()) if workbook_path else None,
        "generated_at": datetime.now(UTC).isoformat(),
    }
    summary.update(stats)
    (output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return summary


async def _run_loop(
    candidates: list[EnrichmentCandidate],
    *,
    limit: int,
    sleep_seconds: float,
    output_dir: Path,
    assigned: int | None = None,
) -> tuple[list[EnrichmentRow], list[QualityIssue]]:
    service = AgentReachService(timeout_seconds=30)
    ready_rows: list[EnrichmentRow] = []
    issues: list[QualityIssue] = []
    checkpoint = output_dir / "agent_reach_checkpoint.jsonl"
    assigned_count = assigned if assigned is not None else min(limit, len(candidates))
    batch = candidates[:limit]

    for index, candidate in enumerate(batch, start=1):
        print(
            f"[{index}/{len(batch)}] {candidate.account_id} {candidate.domain}",
            flush=True,
        )
        rows, issue = await enrich_candidate(candidate, service)
        if rows:
            ready_rows.extend(rows)
            status = "ready"
        elif issue:
            issues.append(issue)
            status = "issue"
        else:
            status = "empty"

        _append_checkpoint(
            checkpoint,
            {
                "account_id": candidate.account_id,
                "company_name": candidate.company_name,
                "domain": candidate.domain,
                "tier": candidate.tier_bucket,
                "status": status,
                "ready_rows": len(rows),
                "issue": issue.issue if issue else None,
                "processed_at": datetime.now(UTC).isoformat(),
            },
        )
        _persist_worker_outputs(
            output_dir,
            ready_rows,
            issues,
            {
                "assigned": assigned_count,
                "attempted": index,
                "incomplete": index < assigned_count,
            },
        )

        if sleep_seconds > 0 and index < len(batch):
            await asyncio.sleep(sleep_seconds)

    return ready_rows, issues


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Run a bounded Agent Reach loop for MUHIDE contact enrichment."
    )
    parser.add_argument(
        "--master-csv",
        type=Path,
        default=_default_downloads_path("MUHIDE_extracted/01_Master_Accounts.csv"),
    )
    parser.add_argument(
        "--existing-xlsx",
        type=Path,
        default=_default_downloads_path("14_Website_Phone_Social_Enrichment_InProgress.xlsx"),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("outputs") / "agent_reach_contact_enrichment" / _timestamp(),
    )
    parser.add_argument("--tiers", default="TIER C", help="Comma-separated tier buckets.")
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Max accounts to process. Default 25 in master mode; all rows for --queue-csv.",
    )
    parser.add_argument("--sleep-seconds", type=float, default=1.0)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument(
        "--no-workbook",
        action="store_true",
        help="Write CSV/checkpoint outputs only; skip the workbook copy.",
    )
    parser.add_argument(
        "--csv-only",
        action="store_true",
        help="Alias for --no-workbook.",
    )
    parser.add_argument(
        "--queue-csv",
        type=Path,
        default=None,
        help="Process this fixed shard/queue instead of scanning the master CSV.",
    )
    return parser


def _resolve_candidates(
    args: argparse.Namespace,
) -> tuple[list[EnrichmentCandidate], set[str], int]:
    """Return (candidates, processed_ids, limit) for master or shard mode."""
    if args.queue_csv is not None:
        candidates = load_queue_candidates(args.queue_csv)
        limit = args.limit if args.limit is not None else len(candidates)
        return candidates, set(), limit

    processed = (
        load_processed_accounts(args.existing_xlsx)
        if args.existing_xlsx.exists()
        else set()
    )
    candidates = load_master_candidates(
        args.master_csv,
        processed_accounts=processed,
        tiers=_parse_tiers(args.tiers),
    )
    limit = args.limit if args.limit is not None else 25
    return candidates, processed, limit


async def async_main() -> int:
    args = _build_parser().parse_args()
    if args.csv_only:
        args.no_workbook = True
    args.output_dir.mkdir(parents=True, exist_ok=True)

    candidates, processed, limit = _resolve_candidates(args)
    assigned = len(candidates) if args.queue_csv is not None else min(limit, len(candidates))
    queue_rows = [queue_row_from_candidate(c) for c in candidates[:limit]]
    _write_csv(
        args.output_dir / "agent_reach_next_queue.csv",
        QUEUE_HEADERS,
        queue_rows,
    )

    print(f"processed_accounts={len(processed)}")
    print(f"remaining_candidates={len(candidates)}")
    print(f"output_dir={args.output_dir.resolve()}")
    if args.queue_csv is not None:
        print(f"queue_csv={args.queue_csv.resolve()}")
        print(f"assigned={len(candidates)} limit={limit}")

    if args.dry_run or limit <= 0:
        print("dry_run=true")
        return 0

    ready_rows, issues = await _run_loop(
        candidates,
        limit=limit,
        sleep_seconds=args.sleep_seconds,
        output_dir=args.output_dir,
        assigned=assigned,
    )

    write_xlsx = worker_writes_workbook(
        queue_csv=args.queue_csv,
        no_workbook=args.no_workbook,
    )
    workbook_path = None
    if write_xlsx and args.existing_xlsx.exists():
        workbook_path = args.output_dir / "14_Website_Phone_Social_Enrichment_AgentReach.xlsx"
        _append_to_workbook(args.existing_xlsx, workbook_path, ready_rows, issues)

    attempted = min(limit, len(candidates))
    summary = _persist_worker_outputs(
        args.output_dir,
        ready_rows,
        issues,
        {
            "queue_csv": str(args.queue_csv.resolve()) if args.queue_csv else None,
            "mode": "worker" if args.queue_csv is not None else "single",
            "assigned": assigned,
            "attempted": attempted,
            "incomplete": attempted < assigned,
        },
        workbook_path=workbook_path,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    return asyncio.run(async_main())


if __name__ == "__main__":
    raise SystemExit(main())
